"""Deterministic DOCX rendering from validated FormatSpec data."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

from format_support import validate_format_spec
from writing_support import (
    collect_global_reference_registry,
    reference_ids_for_node,
    remap_local_citations,
)


_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _flatten_tree(nodes, depth=0):
    for node in nodes if isinstance(nodes, list) else []:
        if not isinstance(node, dict):
            continue
        yield node, depth
        yield from _flatten_tree(node.get("children", []), depth + 1)


def _set_run_font(run, latin, cjk, size_pt=None, bold=None):
    run.font.name = latin
    if size_pt is not None:
        run.font.size = Pt(float(size_pt))
    if bold is not None:
        run.bold = bool(bold)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), cjk)


def _format_paragraph(paragraph, spec, *, indent=True, alignment=None):
    body = spec["body"]
    fmt = paragraph.paragraph_format
    fmt.line_spacing = body["line_spacing"]
    fmt.space_before = Pt(body["space_before_pt"])
    fmt.space_after = Pt(body["space_after_pt"])
    fmt.first_line_indent = Pt(body["font_size_pt"] * body["first_line_indent_chars"] if indent else 0)
    paragraph.alignment = _ALIGNMENTS[alignment or body["alignment"]]
    for run in paragraph.runs:
        _set_run_font(run, body["font_latin"], body["font_cjk"], body["font_size_pt"])


def _set_heading(paragraph, spec, level):
    settings = spec["headings"][f"h{level}"]
    paragraph.alignment = _ALIGNMENTS[settings["alignment"]]
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(settings["space_before_pt"])
    fmt.space_after = Pt(settings["space_after_pt"])
    fmt.first_line_indent = Pt(0)
    for run in paragraph.runs:
        _set_run_font(run, settings["font_latin"], settings["font_cjk"], settings["font_size_pt"], settings["bold"])


def _set_cell_border(cell, **edges):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in edges[edge].items():
            element.set(qn(f"w:{key}"), str(value))


def _apply_three_line_borders(table):
    nil = {"val": "nil"}
    line = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            _set_cell_border(cell, top=nil, bottom=nil, left=nil, right=nil, insideH=nil, insideV=nil)
            if row_index == 0:
                _set_cell_border(cell, top=line)
            if row_index == 0 or row_index == len(table.rows) - 1:
                _set_cell_border(cell, bottom=line if row_index == len(table.rows) - 1 else {"val": "single", "sz": "6", "space": "0", "color": "000000"})


def _clean_inline(value):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", str(value))
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    return re.sub(r"`(.+?)`", r"\1", text).strip()


def _parse_markdown_table(lines):
    rows = []
    for line in lines:
        if line.startswith("TABLE|"):
            line = line[len("TABLE|"):]
        cells = [_clean_inline(cell) for cell in line.strip().strip("|").split("|")]
        if cells and not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
    return rows


def _clean_body_lines(text):
    lines = []
    for raw in str(text).replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            lines.append("")
            continue
        if line.startswith("|") and "|" in line[1:]:
            lines.append("TABLE|" + line)
            continue
        if line in {"[[PAGE_BREAK]]", "\f"}:
            lines.append("[[PAGE_BREAK]]")
            continue
        line = re.sub(r"^#{1,6}\s*", "", line)
        line = re.sub(r"^>+\s*", "", line)
        line = re.sub(r"^[-*+•]\s+", "", line)
        line = re.sub(r"^\d+[.、)]\s*", "", line)
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        line = re.sub(r"\*(.+?)\*", r"\1", line)
        line = re.sub(r"`(.+?)`", r"\1", line)
        lines.append(line)
    return lines


def _add_table(doc, lines, spec, number, locale):
    rows = _parse_markdown_table(lines)
    if not rows:
        return number
    prefix = spec["table"]["caption_prefix"] or ("表" if locale == "zh" else "Table")
    caption = doc.add_paragraph(f"{prefix} {number}")
    _format_paragraph(caption, spec, indent=False, alignment="center")
    for run in caption.runs:
        run.bold = True
        run.font.size = Pt(spec["table"]["font_size_pt"])
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    try:
        table.style = "Table Normal"
    except Exception:
        pass
    for row_index, row in enumerate(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.text = row[col_index] if col_index < len(row) else ""
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    _set_run_font(run, spec["body"]["font_latin"], spec["body"]["font_cjk"], spec["table"]["font_size_pt"], row_index == 0)
    _apply_three_line_borders(table)
    return number + 1


def _add_page_field(section, spec):
    page = spec["page"]["page_number"]
    if not page["enabled"]:
        return
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = _ALIGNMENTS[page["position"]]
    run = paragraph.add_run()
    _set_run_font(run, spec["body"]["font_latin"], spec["body"]["font_cjk"], 10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    displayed = OxmlElement("w:t")
    displayed.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, displayed, end))


def _add_caption(doc, caption, prefix, number, spec):
    label = f"{prefix} {number}"
    text = f"{label}: {caption}" if caption else label
    paragraph = doc.add_paragraph(text)
    _format_paragraph(paragraph, spec, indent=False, alignment="center")
    for run in paragraph.runs:
        run.font.size = Pt(spec["table"]["font_size_pt"])
    return paragraph


def render_manuscript_docx(
    *, tree, drafts, draft_reference_maps, literatures, drafts_charts, drafts_images,
    format_spec, locale="en", template_bytes=None, template_only=False,
):
    """Build an in-memory academic DOCX with explicit FormatSpec precedence over template layout."""
    spec = validate_format_spec(format_spec)
    doc = None
    if template_bytes:
        try:
            doc = Document(io.BytesIO(template_bytes))
        except Exception as exc:
            raise ValueError(f"The selected DOCX template could not be opened: {exc}") from exc
    if doc is None:
        doc = Document()

    for section in doc.sections:
        if not template_only:
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
            margins = spec["page"]["margins_cm"]
            section.top_margin = Cm(margins["top"])
            section.bottom_margin = Cm(margins["bottom"])
            section.left_margin = Cm(margins["left"])
            section.right_margin = Cm(margins["right"])
            _add_page_field(section, spec)

    ordered, number_by_id, literature_by_id = collect_global_reference_registry(
        tree, drafts, draft_reference_maps, literatures
    )
    table_number = 1
    figure_number = 1
    if template_only:
        first_section = doc.sections[0]
        page_width_in = max(1.0, first_section.page_width.inches - first_section.left_margin.inches - first_section.right_margin.inches)
    else:
        page_width_in = max(1.0, 8.2677 - (spec["page"]["margins_cm"]["left"] + spec["page"]["margins_cm"]["right"]) / 2.54)
    for node, depth in _flatten_tree(tree):
        node_id = node.get("id")
        title = str(node.get("title", "") or "")
        if title:
            heading = doc.add_heading(title, level=min(depth + 1, 3))
            if not template_only:
                _set_heading(heading, spec, min(depth + 1, 3))
        text = str(drafts.get(node_id, "") or "")
        if text:
            local_reference_ids = reference_ids_for_node(node, draft_reference_maps)
            text = remap_local_citations(text, local_reference_ids, number_by_id)
            lines = _clean_body_lines(text)
            paragraphs = []
            index = 0
            while index < len(lines):
                line = lines[index]
                if not line or line == "[[PAGE_BREAK]]":
                    if paragraphs:
                        paragraph = doc.add_paragraph("\n".join(paragraphs))
                        if not template_only:
                            _format_paragraph(paragraph, spec)
                        paragraphs = []
                    if line == "[[PAGE_BREAK]]":
                        doc.add_page_break()
                    index += 1
                    continue
                if line.startswith("TABLE|"):
                    if paragraphs:
                        paragraph = doc.add_paragraph("\n".join(paragraphs))
                        if not template_only:
                            _format_paragraph(paragraph, spec)
                        paragraphs = []
                    table_lines = []
                    while index < len(lines) and lines[index].startswith("TABLE|"):
                        table_lines.append(lines[index])
                        index += 1
                    table_number = _add_table(doc, table_lines, spec, table_number, locale)
                    continue
                paragraphs.append(line)
                index += 1
            if paragraphs:
                paragraph = doc.add_paragraph("\n".join(paragraphs))
                if not template_only:
                    _format_paragraph(paragraph, spec)

        for chart in drafts_charts.get(node_id, []) or []:
            path = chart.get("path", "") if isinstance(chart, dict) else ""
            if path:
                try:
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.add_run().add_picture(path, width=Inches(page_width_in))
                    prefix = spec["figure"]["caption_prefix"] or ("图" if locale == "zh" else "Figure")
                    _add_caption(doc, chart.get("caption", "") if isinstance(chart, dict) else "", prefix, figure_number, spec)
                    figure_number += 1
                except (OSError, ValueError):
                    continue

        for image in drafts_images.get(node_id, []) or []:
            path = image.get("path", "") if isinstance(image, dict) else ""
            if path:
                try:
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.add_run().add_picture(path, width=Inches(page_width_in))
                    prefix = spec["figure"]["caption_prefix"] or ("图" if locale == "zh" else "Figure")
                    _add_caption(doc, image.get("caption", "") if isinstance(image, dict) else "", prefix, figure_number, spec)
                    figure_number += 1
                except (OSError, ValueError):
                    continue

    if ordered:
        doc.add_page_break()
        ref_title = "参考文献" if locale == "zh" else "References"
        ref_heading = doc.add_heading(ref_title, level=1)
        if not template_only:
            _set_heading(ref_heading, spec, 1)
        for reference_id in ordered:
            literature = literature_by_id[reference_id]
            title = literature.get("title", "未命名" if locale == "zh" else "Untitled")
            category = literature.get("category", "")
            paragraph = doc.add_paragraph(f"[{number_by_id[reference_id]}] {title} ({category})")
            if not template_only:
                _format_paragraph(paragraph, spec, indent=False, alignment="left")

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
