import io
import json
import re
import unittest
import zipfile
import tempfile
from pathlib import Path

from docx import Document
from PIL import Image

from academic_context_support import attach_chapter_grounding
from format_support import (
    FormatSpecError,
    delete_user_preset,
    load_builtin_presets,
    load_user_presets,
    parse_format_spec,
    save_user_preset,
    validate_format_spec,
)
from document_support import extract_document_text
from literature_intelligence import (
    apply_analysis_to_literature_record,
    analyze_literature_document,
    select_literature_evidence_for_chapter,
    update_evidence_store_after_success,
)
from research_support import (
    PRESET_SECTIONS,
    build_research_planning_context,
    create_default_research_base,
    migrate_research_base,
    normalize_research_source,
    parse_research_source,
    select_research_grounding_for_chapter,
)
from word_export_support import render_manuscript_docx
from writing_support import LLMOutputError, build_chapter_prompt


REPO = Path(__file__).resolve().parents[1]


class ResearchFactLayerTests(unittest.TestCase):
    def test_presets_have_expected_grounding_defaults(self):
        by_id = {item["id"]: item for item in PRESET_SECTIONS}
        self.assertEqual(len(by_id), 8)
        self.assertFalse(by_id["background"]["allow_writing_grounding"])
        self.assertFalse(by_id["requirements"]["allow_writing_grounding"])
        for section_id in ("objectives", "theory", "methods", "data", "results", "constraints"):
            self.assertTrue(by_id[section_id]["allow_writing_grounding"])

    def test_legacy_sections_migrate_without_losing_modules(self):
        legacy = {
            "background": {"modules": [{"title": "B", "content": "background fact"}]},
            "requirements": {"modules": [{"title": "R", "content": "assignment"}]},
            "content": {"modules": [{"title": "Mixed", "content": "keep me"}]},
        }
        migrated = migrate_research_base(legacy)
        sections = migrated["sections"]
        self.assertEqual(sections["background"]["modules"][0]["content"], "background fact")
        self.assertEqual(sections["requirements"]["modules"][0]["content"], "assignment")
        legacy_section = sections["legacy_unclassified"]
        self.assertFalse(legacy_section["allow_writing_grounding"])
        self.assertEqual(legacy_section["modules"][0]["content"], "keep me")

    def test_grounding_off_still_contributes_to_planning_only(self):
        base = create_default_research_base()
        base["sections"]["background"]["modules"] = [{
            "id": "bg1", "title": "Context", "type": "fact", "content": "Historical context",
            "tags": [], "allow_writing_grounding": False,
        }]
        self.assertIn("Historical context", build_research_planning_context(base))
        self.assertNotIn("bg1", [m["id"] for m in select_research_grounding_for_chapter(base, "Introduction", "context")])

    def test_research_parser_processes_complete_ordered_source_and_tail(self):
        source = "FIRST_SENTINEL\n" + ("body sentence. " * 900) + "MIDDLE_SENTINEL\n" + ("more research. " * 1500) + "FINAL_RESEARCH_SENTINEL"
        seen = []

        def llm(prompt, **_kwargs):
            chunk = prompt.split("[DOCUMENT CONTENT]\n", 1)[-1]
            seen.append(chunk)
            markers = [word for word in ("FIRST_SENTINEL", "MIDDLE_SENTINEL", "FINAL_RESEARCH_SENTINEL") if word in chunk]
            content = " ".join(markers) or "General research context"
            return json.dumps({"facts": [{"title": "Extracted fact", "type": "fact", "content": content}]})

        result = parse_research_source(source, "methods", llm, max_chars=2500)
        self.assertEqual(result["status"], "ok")
        self.assertGreater(len(seen), 1)
        self.assertEqual("".join(seen), source)
        combined = " ".join(item["content"] for item in result["facts"])
        for marker in ("FIRST_SENTINEL", "MIDDLE_SENTINEL", "FINAL_RESEARCH_SENTINEL"):
            self.assertIn(marker, combined)
        self.assertEqual([f["source_chunk"] for f in result["facts"]], sorted(f["source_chunk"] for f in result["facts"]))

    def test_failed_parse_returns_no_partial_facts(self):
        source = "A" * 6000
        calls = []

        def llm(prompt, **_kwargs):
            calls.append(prompt)
            if "document part 2/" in prompt.lower():
                raise LLMOutputError("provider unavailable")
            return '[{"title":"Partial","content":"must not persist"}]'

        result = parse_research_source(source, "methods", llm, locale="en", max_chars=2500)
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["facts"], [])

    def test_pasted_and_file_text_share_normalized_source_contract(self):
        pasted = normalize_research_source("study notes", "notes.txt", source_id="same-source")
        uploaded = normalize_research_source(b"study notes", "notes.txt", source_id="same-source")
        self.assertEqual(pasted["text"], uploaded["text"])
        self.assertEqual(pasted["source_id"], uploaded["source_id"])
        self.assertEqual(pasted["text"], "study notes")

    def test_document_text_extractor_reads_full_txt_and_docx_table_order(self):
        txt = "FIRST_SOURCE\n" + ("line of source text.\n" * 1400) + "FINAL_SOURCE_SENTINEL"
        self.assertEqual(extract_document_text(txt.encode("utf-8"), "source.txt"), txt)

        document = Document()
        document.add_paragraph("DOCX_FIRST_SENTINEL")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "TABLE_MIDDLE_SENTINEL"
        table.cell(0, 1).text = "field"
        document.add_paragraph("DOCX_FINAL_SENTINEL")
        raw = io.BytesIO()
        document.save(raw)
        extracted = extract_document_text(raw.getvalue(), "source.docx")
        self.assertLess(extracted.index("DOCX_FIRST_SENTINEL"), extracted.index("TABLE_MIDDLE_SENTINEL"))
        self.assertLess(extracted.index("TABLE_MIDDLE_SENTINEL"), extracted.index("DOCX_FINAL_SENTINEL"))

    def test_research_and_bound_literature_are_attached_to_actual_chapter_prompt(self):
        base = create_default_research_base()
        base["sections"]["methods"]["modules"] = [{
            "id": "fact-1", "section": "methods", "title": "Sampling", "type": "fact",
            "content": "RESEARCH_FACT_SENTINEL", "tags": ["sampling"],
            "allow_writing_grounding": True, "binding_mode": "chapters", "chapter_ids": ["ch-1"],
        }]
        literature = [{
            "id": "lit-1", "title": "Bound Study", "analysis": {"key_findings": "Finding from study"},
        }]
        evidence = {"lit-1": [{
            "literature_id": "lit-1", "chunk_index": 4, "section": "results",
            "claim": "Sampling method", "evidence_text": "BOUND_EVIDENCE_SENTINEL supports the sampling method",
            "source_locator": "chunk 4",
        }]}
        context = attach_chapter_grounding({
            "global_topic": "Research", "target_word_count": 5000,
            "global_outline": "Overall outline", "current_id": "ch-1",
            "current_title": "Sampling", "current_desc": "Describe the method",
            "current_word_count": 250, "current_refs": ["lit-1"],
            "upstream": "Prior section", "downstream": "Later section boundary",
        }, base, evidence, literature, locale="en")
        prompt = build_chapter_prompt(context, {"style_prompt": "Custom style"}, literature, locale="en")
        self.assertIn("RESEARCH_FACT_SENTINEL", prompt)
        self.assertIn("BOUND_EVIDENCE_SENTINEL", prompt)
        self.assertIn("[BOUND LITERATURE EVIDENCE]", prompt)
        self.assertIn("chunk 4", prompt)

    def test_explicit_chapter_binding_precedes_auto_relevance(self):
        base = create_default_research_base()
        section = base["sections"]["methods"]
        section["modules"] = [
            {"id": "bound", "title": "Unrelated wording", "type": "fact", "content": "explicit fact", "tags": [], "allow_writing_grounding": True, "binding_mode": "chapters", "chapter_ids": ["chapter-1"]},
            {"id": "auto", "title": "Sampling methods", "type": "fact", "content": "auto fact", "tags": ["sampling"], "allow_writing_grounding": True, "binding_mode": "auto", "chapter_ids": []},
        ]
        selected = select_research_grounding_for_chapter(base, "Survey sampling", "sampling design", chapter_id="chapter-1")
        self.assertEqual(selected[0]["id"], "bound")
        self.assertIn("auto", [item["id"] for item in selected])


class LiteratureIntelligenceTests(unittest.TestCase):
    @staticmethod
    def _fake_llm(prompt, **_kwargs):
        synthesis = "[DOCUMENT SYNTHESIS]" in prompt
        if synthesis:
            markers = [m for m in ("FIRST_LIT_SENTINEL", "MIDDLE_LIT_SENTINEL", "FINAL_LIT_SENTINEL") if m in prompt]
            return json.dumps({
                "rating": 5, "category": "Methods", "research_question": "RQ",
                "methods": "survey", "sample": "sample", "key_findings": " ".join(markers),
                "limitations": "limits", "quality_assessment": "complete", "summary": " ".join(markers),
            })
        chunk_match = re.search(r"\[SOURCE CHUNK\]\s*(.*?)\s*\[/SOURCE CHUNK\]", prompt, re.S)
        chunk = chunk_match.group(1) if chunk_match else prompt
        markers = [m for m in ("FIRST_LIT_SENTINEL", "MIDDLE_LIT_SENTINEL", "FINAL_LIT_SENTINEL") if m in chunk]
        return json.dumps({
            "research_question": "RQ", "theory": "theory", "method": "survey", "sample": "n=12",
            "data": "survey", "results": "results", "claims": [{"section": "results", "claim": m, "evidence_text": m} for m in markers],
            "limitations": "limits", "conclusion": "conclusion", "definitions": [],
        })

    def test_ten_thousand_character_document_retains_final_evidence(self):
        text = "FIRST_LIT_SENTINEL\n" + ("academic source. " * 300) + "MIDDLE_LIT_SENTINEL\n" + ("paper text. " * 500) + "FINAL_LIT_SENTINEL"
        result = analyze_literature_document(text, "study", self._fake_llm, max_chars=3000)
        self.assertEqual(result["status"], "ok")
        self.assertTrue(result["profile"]["rating"] == 5)
        self.assertIn("FINAL_LIT_SENTINEL", result["profile"]["summary"])
        self.assertIn("FINAL_LIT_SENTINEL", " ".join(e["evidence_text"] for e in result["evidence"]))

    def test_fifty_thousand_character_document_is_chunked_and_synthesized(self):
        text = "FIRST_LIT_SENTINEL\n" + ("long literature. " * 1300) + "MIDDLE_LIT_SENTINEL\n" + ("evidence remains. " * 1500) + "FINAL_LIT_SENTINEL"
        prompts = []
        result = analyze_literature_document(text, "long study", lambda prompt, **kwargs: (prompts.append(prompt) or self._fake_llm(prompt, **kwargs)), max_chars=5000)
        self.assertEqual(result["status"], "ok")
        self.assertGreater(sum("[SOURCE CHUNK " in prompt for prompt in prompts), 2)
        self.assertIn("FIRST_LIT_SENTINEL", result["profile"]["summary"])
        self.assertIn("MIDDLE_LIT_SENTINEL", result["profile"]["summary"])
        self.assertIn("FINAL_LIT_SENTINEL", result["profile"]["summary"])

    def test_failed_chunk_retry_stops_pipeline_and_preserves_prior_evidence(self):
        old = {"lit-1": [{"claim": "KEEP_OLD_EVIDENCE"}]}
        calls = []

        def llm(prompt, **_kwargs):
            calls.append(prompt)
            if "[SOURCE CHUNK 2/" in prompt:
                raise LLMOutputError("offline")
            return self._fake_llm(prompt)

        result = analyze_literature_document("A" * 12000, "study", llm, max_chars=3000)
        updated = update_evidence_store_after_success(old, "lit-1", result)
        self.assertEqual(result["status"], "error")
        self.assertEqual(updated, old)
        self.assertFalse(any("[SOURCE CHUNK 3/" in p for p in calls))

    def test_failed_reanalysis_preserves_literature_profile_and_evidence(self):
        record = {
            "id": "lit-1", "rating": 4, "category": "Methods",
            "analysis": {"summary": "KEEP_OLD_PROFILE"}, "summary": "KEEP_OLD_PROFILE",
        }
        evidence = {"lit-1": [{"claim": "KEEP_OLD_EVIDENCE"}]}
        failed = {"status": "error", "profile": None, "evidence": []}
        updated_record, updated_evidence, success = apply_analysis_to_literature_record(record, evidence, failed)
        self.assertFalse(success)
        self.assertEqual(updated_record, record)
        self.assertEqual(updated_evidence, evidence)

    def test_chapter_evidence_is_limited_to_bound_literature(self):
        evidence = {
            "lit-a": [{"literature_id": "lit-a", "claim": "sampling method", "evidence_text": "random sampling"}],
            "lit-b": [{"literature_id": "lit-b", "claim": "sampling method", "evidence_text": "other study"}],
        }
        selected = select_literature_evidence_for_chapter(evidence, ["lit-a"], "sampling design", limit=5)
        self.assertEqual({item["literature_id"] for item in selected}, {"lit-a"})


class FormatSpecAndWordRendererTests(unittest.TestCase):
    def test_builtin_presets_are_valid_and_include_legacy_compatibility(self):
        presets = load_builtin_presets(REPO / "format_presets")
        self.assertIn("General Academic", presets)
        self.assertIn("Legacy Compatible", presets)
        for spec in presets.values():
            self.assertEqual(validate_format_spec(spec), spec)

    def test_format_spec_rejects_unknown_executable_and_path_fields(self):
        spec = load_builtin_presets(REPO / "format_presets")["General Academic"]
        for key, value in (("python", "print(1)"), ("command", "rm -rf /"), ("path", "../../x")):
            with self.subTest(key=key):
                with self.assertRaises(FormatSpecError):
                    validate_format_spec({**spec, key: value})

    def test_custom_json_roundtrip_and_malformed_input(self):
        spec = load_builtin_presets(REPO / "format_presets")["General Academic"]
        self.assertEqual(parse_format_spec(json.dumps(spec)), spec)
        with self.assertRaises(FormatSpecError):
            parse_format_spec("{bad json")

    def test_custom_preset_store_roundtrip_and_builtin_protection(self):
        builtin = load_builtin_presets(REPO / "format_presets")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "custom.json"
            saved = save_user_preset(path, "My Academic", builtin["General Academic"], builtin_names=builtin)
            self.assertEqual(load_user_presets(path)["My Academic"], saved)
            with self.assertRaises(FormatSpecError):
                save_user_preset(path, "General Academic", builtin["General Academic"], builtin_names=builtin)
            self.assertTrue(delete_user_preset(path, "My Academic", builtin_names=builtin))
            self.assertEqual(load_user_presets(path), {})

    def test_renderer_applies_page_geometry_page_field_and_three_line_tables(self):
        spec = load_builtin_presets(REPO / "format_presets")["General Academic"]
        document = render_manuscript_docx(
            tree=[{"id": "c1", "title": "Methods", "children": []}],
            drafts={"c1": "Intro paragraph.\n\n|Group|N|\n|---|---|\n|A|12|"},
            draft_reference_maps={}, literatures=[], drafts_charts={}, drafts_images={},
            format_spec=spec, locale="en",
        )
        reopened = Document(io.BytesIO(document.getvalue()))
        section = reopened.sections[0]
        self.assertAlmostEqual(section.page_width.inches, 8.27, delta=0.1)
        self.assertGreater(section.top_margin.inches, 0)
        xml = document.getvalue()
        with zipfile.ZipFile(io.BytesIO(xml)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
            footer_xml = "".join(archive.read(name).decode("utf-8") for name in archive.namelist() if re.fullmatch(r"word/footer\d+\.xml", name))
        self.assertIn("w:instrText", footer_xml)
        self.assertIn("PAGE", footer_xml)
        self.assertIn('w:val="nil"', document_xml)
        self.assertIn("Table 1", " ".join(p.text for p in reopened.paragraphs))

    def test_renderer_uses_locale_specific_table_and_figure_captions(self):
        spec = load_builtin_presets(REPO / "format_presets")["General Academic"]
        with tempfile.TemporaryDirectory() as folder:
            image_path = Path(folder) / "figure.png"
            Image.new("RGB", (8, 8), "white").save(image_path)
            for locale, expected in (("en", ("Table 1", "Figure 1: Distribution")), ("zh", ("表 1", "图 1: Distribution"))):
                with self.subTest(locale=locale):
                    output = render_manuscript_docx(
                        tree=[{"id": "c1", "title": "Results", "children": []}],
                        drafts={"c1": "|Group|N|\n|---|---|\n|A|12|"},
                        draft_reference_maps={}, literatures=[], drafts_charts={},
                        drafts_images={"c1": [{"path": str(image_path), "caption": "Distribution"}]},
                        format_spec=spec, locale=locale,
                    )
                    document = Document(io.BytesIO(output.getvalue()))
                    paragraph_text = " ".join(p.text for p in document.paragraphs)
                    for caption in expected:
                        self.assertIn(caption, paragraph_text)

    def test_renderer_preserves_template_content_and_citation_remapping(self):
        template = Document()
        template.add_paragraph("TEMPLATE_CONTENT")
        raw = io.BytesIO()
        template.save(raw)
        output = render_manuscript_docx(
            tree=[{"id": "c1", "title": "Results", "references": ["lit-1"], "children": []}],
            drafts={"c1": "Finding [1]."}, draft_reference_maps={"c1": ["lit-1"]},
            literatures=[{"id": "lit-1", "title": "Evidence source", "category": "Results", "analysis": {"key_findings": "Finding"}}],
            drafts_charts={}, drafts_images={},
            format_spec=load_builtin_presets(REPO / "format_presets")["General Academic"],
            template_bytes=raw.getvalue(), locale="en",
        )
        reopened = Document(io.BytesIO(output.getvalue()))
        text = "\n".join(p.text for p in reopened.paragraphs)
        self.assertIn("TEMPLATE_CONTENT", text)
        self.assertIn("Finding [1]", text)
        self.assertIn("Evidence source", text)


class WiringParityTests(unittest.TestCase):
    def test_both_apps_use_shared_research_literature_and_format_engines(self):
        for app_name in ("app.py", "app_zh.py"):
            source = (REPO / app_name).read_text(encoding="utf-8")
            with self.subTest(app=app_name):
                required = ("research_support", "literature_intelligence", "format_support", "word_export_support")
                missing = [module for module in required if module not in source]
                self.assertFalse(missing, f"{app_name} is not wired to shared modules: {missing}")

    def test_old_fixed_prefix_literature_prompts_are_removed(self):
        for app_name in ("app.py", "app_zh.py"):
            source = (REPO / app_name).read_text(encoding="utf-8")
            self.assertNotRegex(source, r"(?:context|text)\s*\[\s*: ?(?:2000|3000)\s*\]")


if __name__ == "__main__":
    unittest.main()
