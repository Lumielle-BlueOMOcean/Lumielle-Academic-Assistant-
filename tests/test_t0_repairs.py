"""Behavior regressions for the v1.1.0 T0 acceptance findings."""

import copy
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Cm
from PIL import Image

import format_support
import literature_intelligence
import research_support
from academic_context_support import attach_chapter_grounding
from format_ui_support import _generation_prompt
from research_support import (
    build_chapter_research_context,
    build_research_planning_context,
    build_research_planning_prompt_context,
    create_default_research_base,
)
from literature_intelligence import analyze_literature_document, select_literature_evidence_for_chapter
from literature_ui_support import _import_one, render_legacy_reanalysis
from word_export_support import render_manuscript_docx
from writing_support import build_chapter_prompt


class ResearchPlanningDigestTests(unittest.TestCase):
    def test_planning_digest_covers_early_middle_and_late_facts_over_12k(self):
        base = create_default_research_base()
        long = " research detail" * 420
        facts = (
            ("background", "Early", "EARLY_PLANNING_SENTINEL " + long),
            ("requirements", "Middle", long + " MIDDLE_PLANNING_SENTINEL " + long),
            ("results", "Final", long + " FINAL_PLANNING_SENTINEL"),
        )
        for section_id, title, content in facts:
            base["sections"][section_id]["modules"] = [{
                "id": section_id, "title": title, "type": "fact", "content": content,
                "tags": [], "allow_writing_grounding": False,
            }]
        self.assertGreater(sum(len(content) for _, _, content in facts), 12_000)

        digest = build_research_planning_context(base)
        outline_input = build_research_planning_prompt_context(base, "GLOBAL_TOPIC_SENTINEL", locale="zh")

        self.assertLessEqual(len(digest), 12_000)
        for sentinel in ("EARLY_PLANNING_SENTINEL", "MIDDLE_PLANNING_SENTINEL", "FINAL_PLANNING_SENTINEL"):
            self.assertIn(sentinel, digest)
        self.assertIn("Early", digest)
        self.assertIn("Middle", digest)
        self.assertIn("Final", digest)
        self.assertIn("GLOBAL_TOPIC_SENTINEL", outline_input)
        self.assertIn("RESEARCH PLANNING DIGEST", outline_input)
        self.assertIn("FINAL_PLANNING_SENTINEL", outline_input)

    def test_chapter_grounding_keeps_complete_selected_fact_and_tail(self):
        base = create_default_research_base()
        content = "EXPERIMENT_PREFIX " + ("observed sample result; " * 120) + " EXPERIMENT_TAIL_SENTINEL"
        base["sections"]["methods"]["modules"] = [{
            "id": "long-method", "title": "Method detail", "type": "fact", "content": content,
            "tags": ["sampling"], "allow_writing_grounding": True,
        }]

        context = build_chapter_research_context(base, "sampling method", max_chars=8_000)

        self.assertIn("EXPERIMENT_TAIL_SENTINEL", context["research_facts"])
        self.assertIn(content, context["research_facts"])

    def test_chapter_grounding_table_digest_samples_tail_and_marks_omissions(self):
        base = create_default_research_base()
        rows = [[f"row-{index}", f"value-{index}"] for index in range(100)]
        rows[-1][-1] = "TAIL_TABLE_SENTINEL"
        base["sections"]["data"]["modules"] = [{
            "id": "table", "title": "Experiment data", "type": "table", "columns": ["case", "value"],
            "data": rows, "tags": ["data"], "allow_writing_grounding": True,
        }]

        context = build_chapter_research_context(base, "experiment data", max_chars=1_800)

        self.assertIn("TAIL_TABLE_SENTINEL", context["research_facts"])
        self.assertIn("omitted", context["research_facts"].lower())


class LiteratureResearchContextTests(unittest.TestCase):
    def test_research_context_reaches_profile_synthesis_and_changes_relevance_rating(self):
        chunk_prompts = []
        synthesis_prompts = []

        def llm(prompt, **_kwargs):
            if "[DOCUMENT SYNTHESIS]" not in prompt:
                chunk_prompts.append(prompt)
                return json.dumps({
                    "research_question": "A question", "theory": "T", "method": "M",
                    "sample": "S", "data": "D", "results": "R", "claims": [],
                    "limitations": "L", "conclusion": "C", "definitions": [],
                })
            synthesis_prompts.append(prompt)
            relevant = "TOPIC_ALPHA" in prompt
            return json.dumps({
                "rating": 5 if relevant else 1,
                "category": "Methods", "research_question": "A question", "methods": "M",
                "sample": "S", "key_findings": "R", "limitations": "L",
                "quality_assessment": "Strong methods, bounded evidence",
                "relevance_reason": "Directly relevant" if relevant else "Outside this project scope",
                "summary": "Profile",
            })

        def analyze(topic, context):
            try:
                return analyze_literature_document(
                    "A complete source about a method.", "A study", llm,
                    research_topic=topic, research_context=context,
                )
            except TypeError as exc:
                self.fail(f"literature analysis must accept the active research context: {exc}")

        alpha = analyze("TOPIC_ALPHA", "Digest alpha")
        beta = analyze("TOPIC_BETA", "Digest beta")

        self.assertEqual(alpha["status"], "ok")
        self.assertEqual(beta["status"], "ok")
        self.assertGreater(alpha["profile"]["rating"], beta["profile"]["rating"])
        self.assertEqual(alpha["profile"]["relevance_reason"], "Directly relevant")
        self.assertIn("quality_assessment", alpha["profile"])
        self.assertIn("research_question", alpha["profile"])
        self.assertTrue(all("GLOBAL RESEARCH TOPIC" not in prompt for prompt in chunk_prompts))
        self.assertTrue(all("RESEARCH PLANNING CONTEXT" not in prompt for prompt in chunk_prompts))
        self.assertIn("GLOBAL RESEARCH TOPIC", synthesis_prompts[0])
        self.assertIn("TOPIC_ALPHA", synthesis_prompts[0])
        self.assertIn("RESEARCH PLANNING CONTEXT", synthesis_prompts[0])
        self.assertIn("Digest alpha", synthesis_prompts[0])
        self.assertIn("relevance and usefulness", synthesis_prompts[0].lower())
        self.assertIn("quality_assessment", synthesis_prompts[0].lower())

    def test_literature_profile_digest_covers_late_profiles_within_budget(self):
        builder = getattr(literature_intelligence, "build_literature_profile_digest", None)
        self.assertTrue(callable(builder), "literature planning needs a bounded all-profile digest")
        profiles = [
            {"id": "early", "title": "Early source", "rating": 3, "category": "Methods", "analysis": {"key_findings": "EARLY_PROFILE_SENTINEL " + "early evidence " * 250}},
            {"id": "middle", "title": "Middle source", "rating": 3, "category": "Methods", "analysis": {"key_findings": "middle evidence " * 125 + " MIDDLE_PROFILE_SENTINEL " + "middle evidence " * 125}},
            {"id": "final", "title": "Final source", "rating": 3, "category": "Results", "analysis": {"key_findings": "final evidence " * 250 + " FINAL_PROFILE_SENTINEL"}},
        ]

        digest = builder(profiles, max_chars=1_500)

        self.assertLessEqual(len(digest), 1_500)
        for sentinel in ("EARLY_PROFILE_SENTINEL", "MIDDLE_PROFILE_SENTINEL", "FINAL_PROFILE_SENTINEL"):
            self.assertIn(sentinel, digest)
        for title in ("Early source", "Middle source", "Final source"):
            self.assertIn(title, digest)

    def test_literature_library_review_prompt_uses_full_digest_and_context_labels(self):
        builder = getattr(literature_intelligence, "build_literature_library_review_prompt", None)
        self.assertTrue(callable(builder), "overall review needs a shared bounded prompt builder")
        prompt = builder(
            "GLOBAL_TOPIC_SENTINEL", "PLANNING_DIGEST_SENTINEL",
            [{"id": "late", "title": "Late profile", "analysis": {"key_findings": "LATE_LIBRARY_SENTINEL"}}],
        )
        self.assertIn("GLOBAL RESEARCH TOPIC", prompt)
        self.assertIn("GLOBAL_TOPIC_SENTINEL", prompt)
        self.assertIn("RESEARCH PLANNING CONTEXT", prompt)
        self.assertIn("PLANNING_DIGEST_SENTINEL", prompt)
        self.assertIn("LATE_LIBRARY_SENTINEL", prompt)

    def test_oversized_library_review_batches_every_profile_before_final_synthesis(self):
        prompts = []
        records = [
            {"id": f"lit-{index}", "title": f"Source {index}", "rating": 4, "category": "Methods",
             "analysis": {"key_findings": f"PROFILE_{index}_SENTINEL " + ("source detail " * 250)}}
            for index in range(1, 13)
        ]

        def llm(prompt, **_kwargs):
            prompts.append(prompt)
            if "[LITERATURE PROFILE BATCH" in prompt:
                index = next(index for index in range(1, 13) if f"PROFILE_{index}_SENTINEL" in prompt)
                return f"BATCH_SUMMARY_{index}_SENTINEL: " + ("specific source assessment " * 5)
            if "Return a concise intermediate synthesis" in prompt:
                import re
                indices = re.findall(r"BATCH_SUMMARY_(\d+)_SENTINEL", prompt)
                return "BATCH_IDS: " + ",".join(indices)
            self.assertIn("[LITERATURE REVIEW BATCH SUMMARIES]", prompt)
            return "FINAL_LIBRARY_REVIEW_SENTINEL"

        result = literature_intelligence.review_literature_library(
            "GLOBAL_TOPIC", "PLANNING_CONTEXT", records, llm, max_chars=400,
        )

        self.assertEqual(result, "FINAL_LIBRARY_REVIEW_SENTINEL")
        batch_prompts = [prompt for prompt in prompts if "[LITERATURE PROFILE BATCH" in prompt]
        self.assertEqual(len(batch_prompts), 12)
        self.assertTrue(any("Return a concise intermediate synthesis" in prompt for prompt in prompts))
        for index in range(1, 13):
            self.assertTrue(any(f"PROFILE_{index}_SENTINEL" in prompt for prompt in batch_prompts))
        import re
        final_ids = sorted({
            int(value)
            for group in re.findall(r"BATCH_IDS:\s*([0-9,]+)", prompts[-1])
            for value in group.split(",")
        })
        self.assertEqual(final_ids, list(range(1, 13)))
        self.assertIn("GLOBAL_TOPIC", prompts[-1])
        self.assertIn("PLANNING_CONTEXT", prompts[-1])


class FormatPromptSchemaTests(unittest.TestCase):
    def test_natural_language_format_prompt_exposes_all_supported_v1_fields(self):
        prompt = _generation_prompt("Use a compact academic layout.", "en")
        for field in ('"start"', '"keep_with_next"', '"page_break_before"', '"width_percent"', '"caption_enabled"', '"caption_mode"'):
            self.assertIn(field, prompt)
        self.assertIn('"grid"', prompt)
        self.assertIn('"three_line"', prompt)


class LiteratureUiContextWiringTests(unittest.TestCase):
    @staticmethod
    def _llm(prompt, **_kwargs):
        if "[DOCUMENT SYNTHESIS]" in prompt:
            return json.dumps({
                "rating": 4, "category": "Methods", "research_question": "RQ",
                "methods": "M", "sample": "S", "key_findings": "F", "limitations": "L",
                "quality_assessment": "Rigor assessed separately", "relevance_reason": "Project fit",
                "summary": "Summary",
            })
        return json.dumps({
            "research_question": "RQ", "theory": "T", "method": "M", "sample": "S",
            "data": "D", "results": "R", "claims": [], "limitations": "L",
            "conclusion": "C", "definitions": [],
        })

    def test_file_paste_and_legacy_reanalysis_use_project_context(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            records = []
            evidence = {}
            captured = []

            def llm(prompt, **kwargs):
                if "[DOCUMENT SYNTHESIS]" in prompt:
                    captured.append(prompt)
                return self._llm(prompt, **kwargs)

            topic, digest = "CURRENT_TOPIC", "ALL_RESEARCH_FACTS_DIGEST"
            for title, source in (("uploaded.pdf", "Local Upload"), ("pasted.txt", "Pasted Text")):
                evidence, success, _result = _import_one(
                    title, "Complete literature source.", source, "", b"Complete literature source.",
                    Path(title).suffix, records, evidence, root, llm, "en", topic, digest,
                )
                self.assertTrue(success)

            raw_path = root / "legacy.txt"
            raw_path.write_text("Complete legacy source.", encoding="utf-8")
            legacy = {"id": "legacy", "title": "Legacy", "file_path": str(raw_path), "rating": 2}

            class FakeStreamlit:
                @staticmethod
                def caption(_text): pass
                @staticmethod
                def json(_value): pass
                @staticmethod
                def button(_label, **_kwargs): return True
                @staticmethod
                def success(_text): pass
                @staticmethod
                def warning(_text): pass
                @staticmethod
                def rerun(): pass

            render_legacy_reanalysis(
                FakeStreamlit(), legacy, [legacy], evidence, root,
                lambda raw, _filename: raw.decode("utf-8"), llm,
                lambda _records: None, lambda _evidence: None, "en",
                research_topic=topic, research_context=digest,
            )

        self.assertEqual(len(captured), 3)
        for prompt in captured:
            self.assertIn("GLOBAL RESEARCH TOPIC", prompt)
            self.assertIn(topic, prompt)
            self.assertIn("RESEARCH PLANNING CONTEXT", prompt)
            self.assertIn(digest, prompt)


class ResearchLiteratureChapterIntegrationTests(unittest.TestCase):
    def test_research_aware_profile_flows_through_chinese_evidence_into_chapter_prompt(self):
        synthesis_prompts = []
        source = "研究背景 unrelated source. 随机抽样方法按名单抽取样本 RANDOM_SAMPLING_SOURCE_SENTINEL。"

        def llm(prompt, **_kwargs):
            if "[DOCUMENT SYNTHESIS]" in prompt:
                synthesis_prompts.append(prompt)
                return json.dumps({
                    "rating": 5, "category": "Methods", "research_question": "Sampling question",
                    "methods": "Random sampling", "sample": "Study participants",
                    "key_findings": "Sampling method described", "limitations": "Bounded sample",
                    "quality_assessment": "Methods reported; evidence remains source-bounded",
                    "relevance_reason": "Relevant to this project's sampling plan", "summary": "Sampling study",
                })
            return json.dumps({
                "research_question": "Sampling question", "theory": "", "method": "Random sampling",
                "sample": "Study participants", "data": "", "results": "", "limitations": "Bounded sample",
                "conclusion": "", "definitions": [], "claims": [
                    {"section": "Background", "claim": "研究背景", "evidence_text": "研究背景 unrelated source."},
                    {"section": "Methods", "claim": "随机抽样方法", "evidence_text": "随机抽样方法按名单抽取样本 RANDOM_SAMPLING_SOURCE_SENTINEL。"},
                ],
            })

        analysis = analyze_literature_document(
            source, "Sampling study", llm, locale="zh",
            research_topic="GLOBAL_SAMPLING_TOPIC", research_context="PLANNING_SAMPLING_DIGEST",
        )
        self.assertEqual(analysis["status"], "ok")
        self.assertEqual(analysis["profile"]["rating"], 5)
        self.assertIn("GLOBAL_SAMPLING_TOPIC", synthesis_prompts[0])
        self.assertIn("PLANNING_SAMPLING_DIGEST", synthesis_prompts[0])

        record = {"id": "lit-sampling", "title": "Sampling study", "analysis": analysis["profile"]}
        evidence_store = literature_intelligence.update_evidence_store_after_success({}, record["id"], analysis)
        chapter = attach_chapter_grounding(
            {"current_id": "chapter-methods", "current_title": "样本选择与随机抽样方法", "current_desc": "说明抽样流程", "current_refs": [record["id"]]},
            create_default_research_base(), evidence_store, [record], locale="zh",
        )
        prompt = build_chapter_prompt(chapter, {"style_prompt": "STYLE_SENTINEL"}, [record], locale="zh")

        self.assertIn("RANDOM_SAMPLING_SOURCE_SENTINEL", prompt)
        self.assertNotIn("研究背景 unrelated source.", prompt)
        self.assertIn("【本章绑定文献证据】", prompt)


class MultilingualLiteratureEvidenceTests(unittest.TestCase):
    def test_chinese_query_ranks_matching_bound_evidence_and_excludes_unbound(self):
        store = {
            "lit-bound": [
                {"literature_id": "lit-bound", "chunk_index": 1, "claim": "研究背景 unrelated", "evidence_text": "Evidence A"},
                {"literature_id": "lit-bound", "chunk_index": 2, "claim": "随机抽样方法", "evidence_text": "RANDOM_SAMPLING_SENTINEL"},
                {"literature_id": "lit-bound", "chunk_index": 3, "claim": "其他讨论", "evidence_text": "Evidence C"},
            ],
            "lit-unbound": [
                {"literature_id": "lit-unbound", "chunk_index": 1, "claim": "随机抽样方法", "evidence_text": "UNBOUND_SENTINEL"},
            ],
        }

        selected = select_literature_evidence_for_chapter(
            store, ["lit-bound"], "样本选择与随机抽样方法", limit=3,
        )

        self.assertEqual(selected[0]["evidence_text"], "RANDOM_SAMPLING_SENTINEL")
        self.assertTrue(all(item["literature_id"] == "lit-bound" for item in selected))
        self.assertNotIn("UNBOUND_SENTINEL", json.dumps(selected, ensure_ascii=False))

    def test_relevance_ranks_across_multiple_bound_sources(self):
        store = {
            "background": [{"literature_id": "background", "chunk_index": 1, "claim": "研究背景", "evidence_text": "BACKGROUND_EVIDENCE"}],
            "methods": [{"literature_id": "methods", "chunk_index": 1, "claim": "随机抽样方法", "evidence_text": "METHOD_EVIDENCE"}],
        }
        selected = select_literature_evidence_for_chapter(
            store, ["background", "methods"], "随机抽样方法", limit=2,
        )
        self.assertEqual(selected[0]["literature_id"], "methods")

    def test_no_query_terms_uses_explicit_bound_source_order_fallback(self):
        store = {"lit": [
            {"literature_id": "lit", "chunk_index": 2, "claim": "second", "evidence_text": "B"},
            {"literature_id": "lit", "chunk_index": 1, "claim": "first", "evidence_text": "A"},
        ]}

        selected = select_literature_evidence_for_chapter(store, ["lit"], "!!!", limit=2)

        self.assertEqual([item["evidence_text"] for item in selected], ["B", "A"])
        self.assertTrue(all(item["selection_reason"] == "no_query_terms_fallback" for item in selected))


class SmartInboxRoleTests(unittest.TestCase):
    def test_accepted_role_survives_research_base_json_roundtrip(self):
        accept = getattr(research_support, "accept_smart_inbox_facts", None)
        self.assertTrue(callable(accept), "Smart Inbox acceptance needs a persisted role field")
        pending = {"source": {"title": "Study notes"}, "facts": [{"id": "fact-1", "title": "Sampling", "content": "Random sampling"}]}
        accepted = accept(pending, "methods", ["sampling"], "Method evidence", True)
        base = create_default_research_base()
        base["sections"]["methods"]["modules"].extend(accepted)

        restored = research_support.migrate_research_base(json.loads(json.dumps(base)))

        fact = restored["sections"]["methods"]["modules"][0]
        self.assertEqual(fact["role"], "Method evidence")
        self.assertEqual(fact["tags"], ["sampling"])
        self.assertTrue(fact["allow_writing_grounding"])


class FormatSpecV1ContractTests(unittest.TestCase):
    def setUp(self):
        self.presets = format_support.load_builtin_presets(Path(__file__).resolve().parents[1] / "format_presets")

    def test_builtin_presets_have_distinct_general_and_legacy_contracts(self):
        general = self.presets["General Academic"]
        legacy = self.presets["Legacy Compatible"]
        self.assertEqual(general["table"]["style"], "three_line")
        self.assertTrue(general["page"]["page_number"]["enabled"])
        self.assertEqual(general["page"]["margins_cm"]["left"], 3.17)
        self.assertEqual(legacy["table"]["style"], "grid")
        self.assertFalse(legacy["page"]["page_number"]["enabled"])
        self.assertEqual(legacy["page"]["margins_cm"], {"top": 2.54, "bottom": 2.54, "left": 2.54, "right": 2.54})
        self.assertNotEqual(general["table"], legacy["table"])
        self.assertNotEqual(general["page"], legacy["page"])

    def test_grid_page_start_heading_controls_and_figure_width_validate(self):
        spec = copy.deepcopy(self.presets["General Academic"])
        spec["page"]["page_number"]["start"] = 7
        spec["headings"]["h1"]["keep_with_next"] = True
        spec["headings"]["h1"]["page_break_before"] = True
        spec["table"]["style"] = "grid"
        spec["figure"]["width_percent"] = 65

        try:
            clean = format_support.validate_format_spec(spec)
        except format_support.FormatSpecError as exc:
            self.fail(f"FormatSpec v1 must accept the required presentation controls: {exc}")
        self.assertEqual(clean["page"]["page_number"]["start"], 7)
        self.assertTrue(clean["headings"]["h1"]["keep_with_next"])
        self.assertTrue(clean["headings"]["h1"]["page_break_before"])
        self.assertEqual(clean["table"]["style"], "grid")
        self.assertEqual(clean["figure"]["width_percent"], 65)

    def test_page_start_heading_pagination_grid_and_figure_width_render_and_reopen(self):
        spec = copy.deepcopy(self.presets["General Academic"])
        spec["page"]["page_number"]["start"] = 7
        spec["headings"]["h1"]["keep_with_next"] = True
        spec["headings"]["h1"]["page_break_before"] = True
        spec["table"]["style"] = "grid"
        spec["figure"]["width_percent"] = 65
        with tempfile.TemporaryDirectory() as tmp:
            image_path = Path(tmp) / "figure.png"
            Image.new("RGB", (20, 10), "white").save(image_path)
            output = render_manuscript_docx(
                tree=[{"id": "ch", "title": "Heading", "references": ["lit"], "children": []}],
                drafts={"ch": "Body [1]\n|A|B|\n|---|---|\n|x|y|"},
                draft_reference_maps={"ch": ["lit"]},
                literatures=[{"id": "lit", "title": "A Source", "category": "Methods"}],
                drafts_charts={}, drafts_images={"ch": [{"path": str(image_path), "caption": "Figure caption"}]},
                format_spec=spec,
            )
            reopened = Document(output)

        heading = next(p for p in reopened.paragraphs if p.text == "Heading")
        self.assertTrue(heading.paragraph_format.keep_with_next)
        self.assertTrue(heading.paragraph_format.page_break_before)
        self.assertEqual(reopened.tables[0].style.name, "Table Grid")
        self.assertIn("[1] A Source", "\n".join(p.text for p in reopened.paragraphs))
        self.assertEqual(int(reopened.sections[0]._sectPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgNumType").get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}start")), 7)
        available_cm = 21 - spec["page"]["margins_cm"]["left"] - spec["page"]["margins_cm"]["right"]
        self.assertAlmostEqual(reopened.inline_shapes[0].width, Cm(available_cm * .65), delta=1000)

    def test_malformed_page_start_and_figure_width_are_rejected(self):
        for path, value in (("page_start", 0), ("figure_width", 101)):
            spec = copy.deepcopy(self.presets["General Academic"])
            if path == "page_start":
                spec["page"]["page_number"]["start"] = value
            else:
                spec["figure"]["width_percent"] = value
            with self.subTest(path=path), self.assertRaises(format_support.FormatSpecError):
                format_support.validate_format_spec(spec)


if __name__ == "__main__":
    unittest.main()
