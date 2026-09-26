import ast
import io
import importlib.util
import json
import os
import re
import struct
import tempfile
import types
import unittest
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]


def load_project_module(test_case, name):
    path = ROOT / f"{name}.py"
    if not path.is_file():
        test_case.fail(f"Expected project module {path.name} to exist")
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        test_case.fail(f"Could not load project module {path.name}")
    module = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(module)
    except Exception as exc:
        test_case.fail(f"Could not import {path.name}: {exc}")
    return module


def load_app_function(test_case, app_name, function_name, namespace):
    source = (ROOT / app_name).read_text(encoding="utf-8")
    module = ast.parse(source)
    function = next(
        node for node in module.body
        if isinstance(node, ast.FunctionDef) and node.name == function_name
    )
    compiled = compile(ast.Module(body=[function], type_ignores=[]), app_name, "exec")
    function_code = next(
        item for item in compiled.co_consts
        if isinstance(item, types.CodeType) and item.co_name == function_name
    )
    loaded = types.FunctionType(function_code, namespace)
    loaded.__defaults__ = tuple(ast.literal_eval(value) for value in function.args.defaults) or None
    if function.args.kwonlyargs:
        loaded.__kwdefaults__ = {
            argument.arg: ast.literal_eval(value)
            for argument, value in zip(function.args.kwonlyargs, function.args.kw_defaults)
            if value is not None
        }
    return loaded


class CoreCorrectnessTests(unittest.TestCase):
    def test_english_defaults_are_nested_in_prompts_file(self):
        support = load_project_module(self, "writing_support")
        defaults = {
            "prompts.json": {
                "global_topic": "keep this value",
                "aigc_detect_engine": "llm",
            }
        }

        support.apply_english_prompt_defaults(defaults)

        prompts = defaults["prompts.json"]
        self.assertIn("academic paper writing engine", prompts["style_prompt"].lower())
        self.assertIn("page setup", prompts["format_prompt"].lower())
        self.assertIn("rewrite the following academic text", prompts["aigc_rewrite_prompt"].lower())
        self.assertIn("aigc text detection expert", prompts["aigc_detect_prompt"].lower())
        self.assertEqual(prompts["aigc_detect_engine"], "llm")
        self.assertEqual(prompts["global_topic"], "keep this value")
        self.assertEqual(set(defaults), {"prompts.json"})

        app_source = (ROOT / "app.py").read_text(encoding="utf-8")
        self.assertIn("apply_english_prompt_defaults(default_files)", app_source)
        self.assertIsNone(re.search(r'default_files\["(?:style_prompt|format_prompt|aigc_rewrite_prompt|aigc_detect_prompt)"\]', app_source))

    def test_corrupt_user_json_is_preserved_and_subsequent_saves_are_blocked(self):
        class ErrorUI:
            def __init__(self):
                self.errors = []

            def error(self, message):
                self.errors.append(message)

        for app_name in ("app.py", "app_zh.py"):
            with self.subTest(app=app_name), tempfile.TemporaryDirectory() as data_dir:
                filename = "logic_tree.json"
                original_bytes = b'{"title":"KEEP_USER_DATA_SENTINEL",'
                target = Path(data_dir) / filename
                target.write_bytes(original_bytes)
                ui = ErrorUI()
                namespace = {
                    "os": os,
                    "json": json,
                    "DATA_DIR": data_dir,
                    "UNREADABLE_JSON_FILES": set(),
                    "st": ui,
                }
                load_json_file = load_app_function(self, app_name, "load_json_file", namespace)
                save_json_file = load_app_function(self, app_name, "save_json_file", namespace)

                self.assertEqual(load_json_file(filename, []), [])
                self.assertEqual(target.read_bytes(), original_bytes)
                self.assertIn(filename, namespace["UNREADABLE_JSON_FILES"])
                self.assertFalse(save_json_file(filename, []))
                self.assertEqual(target.read_bytes(), original_bytes)
                self.assertTrue(ui.errors)

    def test_mit_open_source_license_and_sidebar_metadata_are_consistent(self):
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        app_sources = {
            name: (ROOT / name).read_text(encoding="utf-8")
            for name in ("app.py", "app_zh.py")
        }
        documents = {"README.md": readme, **app_sources}
        project_text = "\n".join(documents.values())
        for document_name, text in documents.items():
            for required_text in ("MIT License", "1029688024"):
                self.assertIn(required_text, text, document_name)
        self.assertTrue("open-source" in readme.lower() or "开源项目" in readme)
        self.assertIn("https://github.com/Lumielle-BlueOMOcean/Lumielle-Academic-Assistant-", readme)

        restricted_phrases = (
            "仅供个人学习使用",
            "个人学习使用",
            "禁止倒卖",
            "禁止商业牟利",
            "商业牟利",
            "For personal learning use only",
            "Resale for profit is prohibited",
            "commercial profit",
            "non-commercial",
        )
        for phrase in restricted_phrases:
            self.assertNotIn(phrase.lower(), project_text.lower())

        expected_sidebar_text = {
            "app.py": (
                "Open-source project",
                "[MIT License]",
                "[GitHub Repository]",
                "QQ Group **1029688024**",
                "Developer: Lumielle",
            ),
            "app_zh.py": (
                "开源项目",
                "[MIT License]",
                "[GitHub 项目源码]",
                "QQ 群 **1029688024**",
                "开发者：Lumielle",
            ),
        }
        for app_name, required_items in expected_sidebar_text.items():
            module = ast.parse(app_sources[app_name])
            sidebar_function = next(
                node for node in module.body
                if isinstance(node, ast.FunctionDef) and node.name == "render_global_sidebar"
            )
            sidebar_source = ast.get_source_segment(app_sources[app_name], sidebar_function)
            for required_item in required_items:
                self.assertIn(required_item, sidebar_source, app_name)

    def test_context_layers_are_in_chapter_prompts_in_both_locales(self):
        support = load_project_module(self, "writing_support")
        sandwich = {
            "global_topic": "TOPIC_SENTINEL",
            "target_word_count": 900,
            "global_outline": "GLOBAL_SENTINEL",
            "current_title": "CURRENT_SENTINEL",
            "current_desc": "DESC_SENTINEL",
            "current_word_count": 300,
            "current_refs": [],
            "upstream": "UPSTREAM_SENTINEL",
            "downstream": "DOWNSTREAM_SENTINEL",
        }

        for locale in ("en", "zh"):
            with self.subTest(locale=locale):
                prompt = support.build_chapter_prompt(
                    sandwich,
                    {"style_prompt": "STYLE_SENTINEL"},
                    [],
                    locale=locale,
                )
                for sentinel in (
                    "GLOBAL_SENTINEL",
                    "UPSTREAM_SENTINEL",
                    "CURRENT_SENTINEL",
                    "DOWNSTREAM_SENTINEL",
                ):
                    self.assertIn(sentinel, prompt)

    def test_custom_style_prompt_is_interpolated(self):
        support = load_project_module(self, "writing_support")
        sandwich = {
            "global_topic": "Topic",
            "target_word_count": 100,
            "global_outline": "Outline",
            "current_title": "Chapter",
            "current_desc": "Description",
            "current_word_count": 100,
            "current_refs": [],
            "upstream": "None",
            "downstream": "None",
        }

        for locale in ("en", "zh"):
            with self.subTest(locale=locale):
                prompt = support.build_chapter_prompt(
                    sandwich,
                    {"style_prompt": "STYLE_SENTINEL_123"},
                    [],
                    locale=locale,
                )
                self.assertIn("STYLE_SENTINEL_123", prompt)
                self.assertNotIn("{style_prompt}", prompt)

    def test_english_word_counter_counts_words_and_compound_tokens(self):
        support = load_project_module(self, "writing_support")

        self.assertEqual(
            support.count_english_words("Artificial intelligence changes legal research."),
            5,
        )
        self.assertEqual(
            support.count_english_words("OpenAI's state-of-the-art model scored 3.14 in 2024."),
            7,
        )
        self.assertEqual(support.count_english_words("... — !"), 0)

    def test_chinese_character_counter_keeps_non_whitespace_behavior(self):
        support = load_project_module(self, "writing_support")
        self.assertEqual(support.count_chinese_chars("人工 智能，AI"), 7)

    def test_version_is_canonical_v101(self):
        version = load_project_module(self, "version")
        self.assertEqual(version.__version__, "1.0.1")

    def test_readme_matches_canonical_unreleased_version_and_keeps_v100_history(self):
        version = load_project_module(self, "version")
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn(f"v{version.__version__}", readme)
        self.assertRegex(
            readme,
            rf"(?m)^### v{re.escape(version.__version__)} — Unreleased$",
        )
        self.assertIn("v1.0.0", readme)
        self.assertIn("Changelog", readme)

    def test_global_reference_registry_uses_snapshots_and_deduplicates(self):
        support = load_project_module(self, "writing_support")
        tree = [
            {"id": "chapter-a", "references": ["A", "B"], "children": []},
            # Current binding changed after this draft's snapshot was established.
            {"id": "chapter-b", "references": ["A"], "children": []},
        ]
        drafts = {"chapter-a": "A text", "chapter-b": "B text"}
        snapshots = {"chapter-b": ["B", "C"]}
        literatures = [{"id": rid, "title": rid} for rid in ("A", "B", "C")]

        ordered, number_by_id, literature_by_id = support.collect_global_reference_registry(
            tree, drafts, snapshots, literatures
        )

        self.assertEqual(ordered, ["A", "B", "C"])
        self.assertEqual(number_by_id, {"A": 1, "B": 2, "C": 3})
        self.assertEqual(set(literature_by_id), {"A", "B", "C"})

    def test_citations_are_remapped_from_local_to_global_numbers(self):
        support = load_project_module(self, "writing_support")
        # Chapter A binds [A, B], while chapter B binds only [B].
        tree = [
            {"id": "chapter-a", "references": ["A", "B"], "children": []},
            {"id": "chapter-b", "references": ["B"], "children": []},
        ]
        drafts = {"chapter-a": "First", "chapter-b": "Evidence [1]."}
        ordered, number_by_id, _ = support.collect_global_reference_registry(
            tree, drafts, {}, [{"id": "A"}, {"id": "B"}]
        )
        self.assertEqual(ordered, ["A", "B"])
        self.assertEqual(
            support.remap_local_citations(drafts["chapter-b"], ["B"], number_by_id),
            "Evidence [2].",
        )

    def test_multiple_citations_and_unknown_numbers_are_preserved_correctly(self):
        support = load_project_module(self, "writing_support")
        self.assertEqual(
            support.remap_local_citations(
                "Evidence [1][2]. Unknown [99] and year [2024].",
                ["B", "C"],
                {"A": 1, "B": 2, "C": 3},
            ),
            "Evidence [2][3]. Unknown [99] and year [2024].",
        )

    def test_citation_remapping_preserves_markdown_links_and_plain_brackets(self):
        support = load_project_module(self, "writing_support")
        source = "See [1](https://example.test) and [1][source] and [Author, 2024]; cite [1] (as shown).\n[1]: https://example.test/ref"
        self.assertEqual(
            support.remap_local_citations(source, ["B"], {"B": 2}),
            "See [1](https://example.test) and [1][source] and [Author, 2024]; cite [2] (as shown).\n[1]: https://example.test/ref",
        )

    def test_legacy_reference_registry_falls_back_to_tree_bindings(self):
        support = load_project_module(self, "writing_support")
        tree = [{"id": "legacy", "references": ["B", "A", "B"], "children": []}]
        ordered, numbers, _ = support.collect_global_reference_registry(
            tree,
            {"legacy": "Legacy draft [1]."},
            {},
            [{"id": "A"}, {"id": "B"}],
        )
        self.assertEqual(ordered, ["B", "A"])
        self.assertEqual(numbers, {"B": 1, "A": 2})

    def test_llm_failures_are_rejected_but_normal_prose_is_accepted(self):
        support = load_project_module(self, "writing_support")
        for failure in (
            None,
            "",
            "   \n\t",
            "API key missing. Add a key.",
            "API call failed repeatedly: timeout",
            "OpenAI client initialization error: bad config",
            "No LLM configured. Please add one.",
            "No valid LLM configuration found.",
            "API 密钥缺失，请在基础配置中补全。",
            "API 连续调用失败: timeout",
            "OpenAI 客户端构造异常: invalid URL",
            "未配置任何 LLM，请先添加。",
            "未找到有效的 LLM 配置。",
        ):
            with self.subTest(failure=failure):
                self.assertTrue(support.get_llm_failure_reason(failure))
                with self.assertRaises(support.LLMOutputError):
                    support.require_valid_llm_output(failure)

        prose = "The study compares API access policies across three research groups."
        self.assertIsNone(support.get_llm_failure_reason(prose))
        self.assertEqual(support.require_valid_llm_output(prose), prose)
        quoted_error = "The paper analyzes the interface message 'API key missing' as an example."
        self.assertIsNone(support.get_llm_failure_reason(quoted_error))

    def test_failed_generation_does_not_change_draft_or_reference_snapshot(self):
        support = load_project_module(self, "writing_support")
        drafts = {"chapter": "Old body"}
        snapshots = {"chapter": ["old-reference"]}
        with self.assertRaises(support.LLMOutputError):
            support.record_generated_draft(drafts, snapshots, "chapter", "", ["new-reference"])
        self.assertEqual(drafts, {"chapter": "Old body"})
        self.assertEqual(snapshots, {"chapter": ["old-reference"]})

    def test_dispatch_returns_valid_text_or_raises_for_known_failures(self):
        support = load_project_module(self, "writing_support")
        for app_name in ("app.py", "app_zh.py"):
            namespace = {
                "get_all_profiles": lambda: [{"id": "profile", "api_key": "key", "base_url": "url", "model": "model"}],
                "get_active_profile": lambda: {"id": "profile", "api_key": "key", "base_url": "url", "model": "model"},
                "call_llm_api": lambda **kwargs: "valid model output",
                "checked_llm_call": support.checked_llm_call,
                "require_valid_llm_output": support.require_valid_llm_output,
                "LLMOutputError": support.LLMOutputError,
            }
            dispatch = load_app_function(self, app_name, "dispatch_llm_call", namespace)
            self.assertEqual(dispatch("prompt"), "valid model output", app_name)
            namespace["get_all_profiles"] = lambda: []
            with self.assertRaises(support.LLMOutputError):
                dispatch("prompt")
            namespace["get_all_profiles"] = lambda: [{"id": "profile", "api_key": "key", "base_url": "url", "model": "model"}]

            for response in (
                "API key missing: sk-test-secret-token",
                "API call failed repeatedly: timeout",
                "OpenAI client initialization error: bad config",
                "No LLM configured.",
                "No valid LLM configuration found.",
                "API 密钥缺失：sk-test-secret-token",
                "API 连续调用失败：timeout",
                "OpenAI 客户端构造异常：bad config",
                "未配置任何 LLM。",
                "未找到有效的 LLM 配置。",
                "",
                "   ",
                None,
            ):
                with self.subTest(app=app_name, response=response):
                    namespace["call_llm_api"] = lambda **kwargs: response
                    with self.assertRaises(support.LLMOutputError):
                        dispatch("prompt")

    def test_partial_model_discussion_keeps_failures_out_of_answers(self):
        support = load_project_module(self, "writing_support")

        def fake_model(model_id):
            if model_id == "B":
                raise support.LLMOutputError("API call failed repeatedly: timeout")
            return f"answer from {model_id}"

        answers, failures = support.collect_model_answers(("A", "B", "C"), fake_model)
        self.assertEqual(answers, {"A": "answer from A", "C": "answer from C"})
        self.assertEqual(failures, ["B"])

    def test_literature_analysis_failure_is_unrated_but_success_is_marked(self):
        support = load_project_module(self, "writing_support")
        for response in ("API call failed repeatedly: timeout", "not JSON"):
            with self.subTest(response=response):
                result = support.parse_literature_analysis(response)
                self.assertEqual(result["rating"], 0)
                self.assertEqual(result["category"], "Unrated")
                self.assertEqual(result["analysis_status"], "unavailable")

        result = support.parse_literature_analysis(
            '{"rating":4,"category":"Methods","key_findings":"Finding"}'
        )
        self.assertEqual(result["rating"], 4)
        self.assertEqual(result["category"], "Methods")
        self.assertEqual(result["analysis_status"], "ok")

    def test_aigc_failure_has_no_probability(self):
        support = load_project_module(self, "writing_support")
        for locale in ("en", "zh"):
            with self.subTest(locale=locale):
                result = support.aigc_detection_error_result(locale)
                self.assertIsNone(result["score"])
                self.assertEqual(result["label"], "Error")

        for app_name in ("app.py", "app_zh.py"):
            def fail_dispatch(*args, **kwargs):
                raise support.LLMOutputError("API call failed repeatedly: sk-test-secret-token")

            namespace = {
                "load_json_file": lambda *args, **kwargs: {},
                "dispatch_llm_call": fail_dispatch,
                "aigc_detection_error_result": support.aigc_detection_error_result,
                "LLMOutputError": support.LLMOutputError,
                "re": re,
                "json": json,
            }
            detect = load_app_function(self, app_name, "llm_detect_aigc", namespace)
            result = detect("a sufficiently long sample text")
            self.assertIsNone(result["score"], app_name)
            self.assertEqual(result["label"], "Error", app_name)

    def test_deep_review_report_uses_successful_blocks_only(self):
        support = load_project_module(self, "writing_support")
        report, successful, failed = support.build_deep_review_report(
            [
                {"status": "ok", "title": "Chapter A", "content": "Review A"},
                {"status": "error", "title": "Chapter B", "content": "API key missing: sk-secret"},
                {"status": "ok", "title": "Chapter C", "content": "Review C"},
            ],
            locale="en",
        )
        self.assertEqual((successful, failed), (2, 1))
        self.assertIn("Review A", report)
        self.assertIn("Review C", report)
        self.assertNotIn("sk-secret", report)
        self.assertIsNone(support.build_deep_review_report([], locale="en")[0])

    def test_apps_use_centralized_failure_boundary_and_user_safe_messages(self):
        for app_name in ("app.py", "app_zh.py"):
            source = (ROOT / app_name).read_text(encoding="utf-8")
            with self.subTest(app=app_name):
                self.assertNotRegex(
                    source,
                    r"\.startswith\(\s*['\"](?:API key missing|API call failed repeatedly|API 密钥缺失|API 连续调用失败)",
                )
                dispatch_body = source.split("def dispatch_llm_call(", 1)[1].split("\ndef ", 1)[0]
                self.assertIn("checked_llm_call", dispatch_body)
                self.assertIn("llm_error_message", source)
                self.assertIn("collect_model_answers", source)
                self.assertIn("parse_literature_analysis", source)
                self.assertIn("build_deep_review_report", source)
                self.assertNotIn("{exc}", source)
                review_body = source.split("def render_full_logic_review(", 1)[1].split("\ndef module5_writing", 1)[0]
                self.assertLess(
                    review_body.index("if final_report is None:"),
                    review_body.index('save_json_file("full_logic_review.json"'),
                )

        support = load_project_module(self, "writing_support")
        failure = support.get_llm_failure_reason(
            "API call failed repeatedly: invalid API key sk-test-secret-token"
        )
        self.assertNotIn("sk-test-secret-token", failure)


class DocumentSupportTests(unittest.TestCase):
    def setUp(self):
        self.support = load_project_module(self, "document_support")

    def test_chunks_cover_long_document_without_loss_or_overlap(self):
        source = ("Paragraph one.\n\n" + "x" * 6200 + "\n\nParagraph three.\n") * 2
        chunks = self.support.chunk_document_text(source, max_chars=5500)

        self.assertGreater(len(chunks), 2)
        self.assertEqual("".join(chunks), source)
        self.assertTrue(all(0 < len(chunk) <= 5500 for chunk in chunks))

    def test_short_and_empty_documents_have_predictable_chunks(self):
        self.assertEqual(self.support.chunk_document_text("short text"), ["short text"])
        self.assertEqual(self.support.chunk_document_text(""), [])
        whitespace = "  \n\t"
        self.assertEqual("".join(self.support.chunk_document_text(whitespace)), whitespace)

    def test_parse_error_detection_covers_both_locales(self):
        for error in (
            "Parse failed: unreadable file",
            "  Parse failed: scanned PDF",
            "解析失败: 文件无法读取",
        ):
            with self.subTest(error=error):
                self.assertTrue(self.support.is_document_parse_error(error))
        self.assertFalse(self.support.is_document_parse_error("The paper starts with Parse failed as a phrase."))
        self.assertFalse(self.support.is_document_parse_error("正常提取的研究内容"))

    def test_all_chunks_reach_llm_in_order_and_last_sentinel_is_not_truncated(self):
        source = "A" * 6100 + "\n\nMIDDLE_SENTINEL\n\n" + "B" * 6100 + "\n\nEND_SENTINEL_456"
        prompts = []

        def llm_call(prompt, **kwargs):
            prompts.append((prompt, kwargs))
            part = re.search(r"DOCUMENT PART (\d+)/(\d+)", prompt)
            self.assertIsNotNone(part)
            return json.dumps([{"title": f"Part {part.group(1)}", "content": f"body-{part.group(1)}"}])

        result = self.support.parse_document_in_chunks(
            source,
            "research content",
            llm_call,
            parse_response=json.loads,
            normalize_modules=lambda values: values,
            locale="en",
        )

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["chunks_total"], 3)
        self.assertEqual([module["title"] for module in result["modules"]], [f"Part {i}" for i in range(1, 4)])
        self.assertEqual(len(prompts), 3)
        self.assertIn("END_SENTINEL_456", prompts[-1][0])
        self.assertTrue(all(call[1]["max_tokens"] <= 8000 for call in prompts))

    def test_chinese_prompt_identifies_part_and_forbids_inference(self):
        prompts = []

        def llm_call(prompt, **kwargs):
            prompts.append(prompt)
            return '[{"title":"研究要求","content":"只包含本块内容"}]'

        result = self.support.parse_document_in_chunks(
            "中文文档末尾内容",
            "研究要求",
            llm_call,
            parse_response=json.loads,
            normalize_modules=lambda values: values,
            locale="zh",
        )

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["modules"][0]["title"], "研究要求")
        self.assertEqual(len(prompts), 1)
        self.assertIn("DOCUMENT PART 1/1", prompts[0])
        self.assertIn("不得推断", prompts[0])
        self.assertIn("中文文档末尾内容", prompts[0])

    def test_retries_only_the_failed_chunk_once_and_preserves_module_order(self):
        source = "A" * 5500 + "\n\n" + "B" * 5500 + "\n\n" + "C" * 100
        calls = []
        prompts = []
        attempts_by_part = {}

        def llm_call(prompt, **kwargs):
            part = int(re.search(r"DOCUMENT PART (\d+)/(\d+)", prompt).group(1))
            calls.append(part)
            prompts.append(prompt)
            attempts_by_part[part] = attempts_by_part.get(part, 0) + 1
            if part == 2 and attempts_by_part[part] == 1:
                return "not json"
            return json.dumps([{"title": f"Part {part}", "content": str(part)}])

        result = self.support.parse_document_in_chunks(
            source,
            "background",
            llm_call,
            parse_response=json.loads,
            normalize_modules=lambda values: values,
            locale="en",
        )

        self.assertEqual(calls, [1, 2, 2, 3])
        self.assertEqual(prompts[1], prompts[2].replace("Your previous response was not a valid non-empty module JSON array. Correct it and output only valid JSON.\n", ""))
        self.assertEqual(result["status"], "ok")
        self.assertEqual([module["title"] for module in result["modules"]], ["Part 1", "Part 2", "Part 3"])

    def test_exhausted_chunk_failure_returns_no_partial_modules(self):
        source = "A" * 5500 + "\n\n" + "B" * 5500 + "\n\n" + "C" * 100
        calls = []
        logged_failures = []

        def llm_call(prompt, **kwargs):
            part = int(re.search(r"DOCUMENT PART (\d+)/(\d+)", prompt).group(1))
            calls.append(part)
            return json.dumps([{"title": "first"}]) if part == 1 else "not json"

        result = self.support.parse_document_in_chunks(
            source,
            "requirements",
            llm_call,
            parse_response=lambda raw: json.loads(raw) if raw.startswith("[") else None,
            normalize_modules=lambda values: values,
            on_parse_failure=lambda chunk, response: logged_failures.append((chunk, response)),
            locale="en",
        )

        self.assertEqual(calls, [1, 2, 2])
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["modules"], [])
        self.assertEqual(result["error_type"], "parse")
        self.assertIn("no partial content was imported", result["message"])
        self.assertEqual(len(logged_failures), 1)
        self.assertTrue(logged_failures[0][0].lstrip().startswith("B"))

    def test_llm_failure_returns_no_partial_modules_and_does_not_retry(self):
        source = "A" * 5500 + "\n\n" + "B" * 100
        calls = []

        def llm_call(prompt, **kwargs):
            part = int(re.search(r"DOCUMENT PART (\d+)/(\d+)", prompt).group(1))
            calls.append(part)
            if part == 1:
                return json.dumps([{"title": "first"}])
            return "API call failed repeatedly: timeout"

        result = self.support.parse_document_in_chunks(
            source,
            "requirements",
            llm_call,
            parse_response=json.loads,
            normalize_modules=lambda values: values,
            locale="en",
        )

        self.assertEqual(calls, [1, 2])
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["modules"], [])
        self.assertEqual(result["error_type"], "llm")
        self.assertIn("no partial content was imported", result["message"])

    def test_applications_delegate_full_document_and_use_shared_parse_error_check(self):
        for app_name in ("app.py", "app_zh.py"):
            source = (ROOT / app_name).read_text(encoding="utf-8")
            with self.subTest(app=app_name):
                self.assertIn("parse_document_in_chunks(", source)
                self.assertIn('parse_result["status"] == "ok"', source)
                self.assertGreaterEqual(source.count("is_document_parse_error("), 2)
                self.assertNotIn("doc_text[:6000]", source)
                self.assertNotIn("doc_text[:5000]", source)


class OutlineSupportTests(unittest.TestCase):
    def setUp(self):
        self.support = load_project_module(self, "outline_support")
        self.documents = load_project_module(self, "document_support")

    @staticmethod
    def tree_response(title="Final outline", **fields):
        node = {"title": title, "desc": "Source structure", "children": []}
        node.update(fields)
        return json.dumps([node], ensure_ascii=False)

    def test_short_manuscript_uses_one_final_tree_call_with_all_text(self):
        source = "Introduction sentinel\nMethods sentinel\nConclusion sentinel"
        prompts = []

        def llm_call(prompt, **kwargs):
            prompts.append(prompt)
            return self.tree_response("Paper", word_count="24")

        result = self.support.reverse_engineer_manuscript(source, llm_call, locale="en")

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["chunks_total"], 1)
        self.assertEqual(len(prompts), 1)
        self.assertNotIn("STAGE 1: SOURCE STRUCTURE EXTRACTION", prompts[0])
        for sentinel in ("Introduction sentinel", "Methods sentinel", "Conclusion sentinel"):
            self.assertIn(sentinel, prompts[0])
        self.assertEqual(result["tree"][0]["word_count"], 24)

    def test_long_manuscript_processes_each_chunk_once_and_keeps_final_sentinel(self):
        source = "BEGIN_SENTINEL\n" + ("body text " * 2200) + "\nMIDDLE_SENTINEL\n" + ("methods text " * 2200) + "\nFINAL_SECTION_SENTINEL"
        expected_chunks = self.documents.chunk_document_text(source)
        stage_one_prompts = []

        def llm_call(prompt, **kwargs):
            if "STAGE 1: SOURCE STRUCTURE EXTRACTION" in prompt:
                stage_one_prompts.append(prompt)
                match = re.search(r"SOURCE DOCUMENT PART (\d+)/(\d+)", prompt)
                return self.tree_response(f"Part {match.group(1)}")
            return self.tree_response("Consolidated manuscript")

        result = self.support.reverse_engineer_manuscript(source, llm_call, locale="en")

        self.assertGreater(len(source), 20000)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["chunks_total"], len(expected_chunks))
        self.assertEqual(len(stage_one_prompts), len(expected_chunks))
        self.assertEqual(
            [int(re.search(r"SOURCE DOCUMENT PART (\d+)/(\d+)", prompt).group(1)) for prompt in stage_one_prompts],
            list(range(1, len(expected_chunks) + 1)),
        )
        self.assertIn("FINAL_SECTION_SENTINEL", stage_one_prompts[-1])

    def test_long_manuscript_retries_only_invalid_source_chunk(self):
        source = "A" * 5500 + "\n\n" + "B" * 5500 + "\n\n" + "C" * 100
        calls = []
        attempts = {}

        def llm_call(prompt, **kwargs):
            if "STAGE 1: SOURCE STRUCTURE EXTRACTION" not in prompt:
                return self.tree_response("Consolidated")
            part = int(re.search(r"SOURCE DOCUMENT PART (\d+)/(\d+)", prompt).group(1))
            calls.append(part)
            attempts[part] = attempts.get(part, 0) + 1
            if part == 2 and attempts[part] == 1:
                return "not JSON"
            return self.tree_response(f"Part {part}")

        result = self.support.reverse_engineer_manuscript(source, llm_call, locale="en")

        self.assertEqual(calls, [1, 2, 2, 3])
        self.assertEqual(result["status"], "ok")

    def test_ordered_consolidation_merges_cross_chunk_chapters(self):
        source = "A" * 5500 + "\n\n" + "B" * 5500 + "\n\n" + "C" * 100
        merge_prompts = []
        partials = {
            1: [{"title": "Methods", "children": [{"title": "Dataset", "children": []}]}],
            2: [{"title": "Methods continuation", "children": [{"title": "Model", "children": []}]}],
            3: [{"title": "Results", "children": []}],
        }
        unified = [
            {"title": "Methods", "children": [
                {"title": "Dataset", "children": []}, {"title": "Model", "children": []}
            ]},
            {"title": "Results", "children": []},
        ]

        def llm_call(prompt, **kwargs):
            if "STAGE 1: SOURCE STRUCTURE EXTRACTION" in prompt:
                part = int(re.search(r"SOURCE DOCUMENT PART (\d+)/(\d+)", prompt).group(1))
                return json.dumps(partials[part])
            merge_prompts.append(prompt)
            return json.dumps(unified)

        result = self.support.reverse_engineer_manuscript(source, llm_call, locale="en")

        self.assertEqual(result["status"], "ok")
        self.assertEqual([node["title"] for node in result["tree"]], ["Methods", "Results"])
        self.assertEqual([node["title"] for node in result["tree"][0]["children"]], ["Dataset", "Model"])
        self.assertEqual(len(merge_prompts), 1)
        positions = [merge_prompts[0].index(label) for label in ("Methods", "Dataset", "Methods continuation", "Model", "Results")]
        self.assertEqual(positions, sorted(positions))

    def test_large_consolidation_batches_all_fragments_with_bounded_progress(self):
        source = "x" * (5500 * 8)
        merge_prompts = []

        def llm_call(prompt, **kwargs):
            if "STAGE 1: SOURCE STRUCTURE EXTRACTION" in prompt:
                part = int(re.search(r"SOURCE DOCUMENT PART (\d+)/(\d+)", prompt).group(1))
                return json.dumps([{"title": f"FRAGMENT_{part}", "desc": "d" * 4500, "children": []}])
            merge_prompts.append(prompt)
            return self.tree_response("Merged fragment")

        result = self.support.reverse_engineer_manuscript(source, llm_call, locale="en")

        self.assertEqual(result["status"], "ok")
        self.assertGreater(result["merge_rounds"], 1)
        self.assertTrue(merge_prompts)
        self.assertTrue(all(len(prompt) <= self.support.MAX_CONSOLIDATION_PROMPT_CHARS for prompt in merge_prompts))
        round_one = [prompt for prompt in merge_prompts if "CONSOLIDATION ROUND 1" in prompt]
        self.assertGreater(len(round_one), 1)
        for part in range(1, 9):
            self.assertEqual(sum(prompt.count(f"FRAGMENT_{part}") for prompt in round_one), 1)

    def test_nonprogressing_or_oversized_merge_fails_without_tree(self):
        source = "A" * 5500 + "\n\n" + "B" * 100
        prompts = []

        def llm_call(prompt, **kwargs):
            prompts.append(prompt)
            if "STAGE 1: SOURCE STRUCTURE EXTRACTION" in prompt:
                part = int(re.search(r"SOURCE DOCUMENT PART (\d+)/(\d+)", prompt).group(1))
                return json.dumps([{"title": f"Part {part}", "desc": "x" * 26000}])
            return self.tree_response("Unexpected")

        result = self.support.reverse_engineer_manuscript(source, llm_call, locale="en")

        self.assertEqual(result["status"], "error")
        self.assertNotIn("tree", result)
        self.assertFalse(any("STAGE 2: ORDERED CONSOLIDATION" in prompt for prompt in prompts))

    def test_normalization_rebuilds_ids_fields_and_parent_word_counts(self):
        normalized = self.support.normalize_reverse_outline_tree([
            {"id": "same-id", "title": " Parent ", "word_count": "999", "references": ["fake"],
             "is_own_experiment": "false", "children": [
                 {"id": "same-id", "title": " Child A ", "word_count": "12"},
                 {"id": "same-id", "title": "Child B", "word_count": -4, "children": "invalid"},
                 None,
                 {"title": "   ", "word_count": 90},
             ]},
            {"title": "Leaf", "word_count": "7", "children": None},
        ])

        ids = [node["id"] for node in normalized]
        ids.extend(child["id"] for node in normalized for child in node["children"])
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(normalized[0]["title"], "Parent")
        self.assertEqual(normalized[0]["word_count"], 12)
        self.assertEqual(normalized[0]["references"], [])
        self.assertEqual(normalized[0]["is_own_experiment"], False)
        self.assertEqual(normalized[0]["chart_instruction"], "")
        self.assertEqual(normalized[0]["image_suggestion"], "")
        self.assertEqual(normalized[0]["children"][1]["word_count"], 0)
        self.assertEqual(normalized[1]["word_count"], 7)

    def test_docx_extraction_includes_table_paragraphs_in_document_order(self):
        from docx import Document

        document = Document()
        document.add_paragraph("INTRODUCTION_SENTINEL")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "TABLE_SECTION_SENTINEL"
        document.add_paragraph("CONCLUSION_SENTINEL")

        extracted = self.support.extract_docx_manuscript_text(document)

        self.assertLess(extracted.index("INTRODUCTION_SENTINEL"), extracted.index("TABLE_SECTION_SENTINEL"))
        self.assertLess(extracted.index("TABLE_SECTION_SENTINEL"), extracted.index("CONCLUSION_SENTINEL"))

    def test_duplicate_model_ids_are_replaced_with_unique_local_ids(self):
        source = "short source"

        def llm_call(prompt, **kwargs):
            return json.dumps([
                {"id": "same-id", "title": "One", "children": []},
                {"id": "same-id", "title": "Two", "children": []},
            ])

        result = self.support.reverse_engineer_manuscript(source, llm_call)
        ids = [node["id"] for node in result["tree"]]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertNotEqual(ids, ["same-id", "same-id"])

    def test_llm_failure_never_provides_persistable_partial_tree(self):
        source = "A" * 5500 + "\n\n" + "B" * 100

        def llm_call(prompt, **kwargs):
            if "STAGE 1: SOURCE STRUCTURE EXTRACTION" in prompt and "SOURCE DOCUMENT PART 2/2" in prompt:
                raise self.support.LLMOutputError("API unavailable")
            return self.tree_response("Partial")

        result = self.support.reverse_engineer_manuscript(source, llm_call)
        self.assertEqual(result["status"], "error")
        self.assertNotIn("tree", result)
        with tempfile.TemporaryDirectory() as data_dir:
            self.assertFalse(self.support.persist_reverse_outline_result(result, data_dir))
            self.assertFalse((Path(data_dir) / "logic_tree.json").exists())

    def test_invalid_final_json_is_retried_once_then_fails_closed(self):
        calls = []

        def llm_call(prompt, **kwargs):
            calls.append(prompt)
            return "not JSON"

        result = self.support.reverse_engineer_manuscript("short source", llm_call)
        self.assertEqual(result["status"], "error")
        self.assertNotIn("tree", result)
        self.assertEqual(len(calls), 2)
        with tempfile.TemporaryDirectory() as data_dir:
            self.assertFalse(self.support.persist_reverse_outline_result(result, data_dir))
            self.assertFalse((Path(data_dir) / "logic_tree.json").exists())

    def test_overdeep_final_outline_gets_one_correction_retry(self):
        prompts = []

        def llm_call(prompt, **kwargs):
            prompts.append(prompt)
            if len(prompts) == 1:
                return json.dumps([{"title": "One", "children": [{"title": "Two", "children": [{"title": "Three", "children": [{"title": "Too deep"}]}]}]}])
            return json.dumps([{"title": "One", "children": [{"title": "Two", "children": [{"title": "Three"}]}]}])

        result = self.support.reverse_engineer_manuscript("short source", llm_call)

        self.assertEqual(result["status"], "ok")
        self.assertEqual(len(prompts), 2)
        self.assertIn("previous result", prompts[1].lower())
        self.assertEqual(len(result["tree"][0]["children"][0]["children"][0]["children"]), 0)

    def test_empty_manuscript_returns_before_any_llm_call(self):
        calls = []
        result = self.support.reverse_engineer_manuscript(" \n\t", lambda *args, **kwargs: calls.append(args))
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["error_type"], "empty")
        self.assertNotIn("tree", result)
        self.assertEqual(calls, [])

    def test_consolidation_llm_failure_retries_group_then_returns_no_tree(self):
        source = "A" * 5500 + "\n\n" + "B" * 100
        merge_calls = []

        def llm_call(prompt, **kwargs):
            if "STAGE 1: SOURCE STRUCTURE EXTRACTION" in prompt:
                return self.tree_response("Local")
            merge_calls.append(prompt)
            raise self.support.LLMOutputError("API unavailable")

        result = self.support.reverse_engineer_manuscript(source, llm_call)

        self.assertEqual(len(merge_calls), 2)
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["stage"], "consolidation")
        self.assertNotIn("tree", result)

    def test_persistence_gate_saves_only_successful_nonempty_tree(self):
        tree = [{"id": "local", "title": "Complete", "children": []}]
        with tempfile.TemporaryDirectory() as data_dir:
            target = Path(data_dir) / "logic_tree.json"
            previous_tree = '[{"id":"previous","title":"Keep"}]'
            target.write_text(previous_tree, encoding="utf-8")
            self.assertFalse(self.support.persist_reverse_outline_result({"status": "error", "tree": [{"title": "partial"}]}, data_dir))
            self.assertFalse(self.support.persist_reverse_outline_result({"status": "ok", "tree": []}, data_dir))
            self.assertEqual(target.read_text(encoding="utf-8"), previous_tree)
            self.assertTrue(self.support.persist_reverse_outline_result({"status": "ok", "tree": tree}, data_dir))
            self.assertEqual(json.loads(target.read_text(encoding="utf-8")), tree)

    def test_atomic_persistence_keeps_old_tree_when_temp_write_fails(self):
        tree = [{"id": "new", "title": "Complete", "children": []}]
        with tempfile.TemporaryDirectory() as data_dir:
            target = Path(data_dir) / "logic_tree.json"
            target.write_text('[{"id":"old","title":"Keep"}]', encoding="utf-8")
            old_bytes = target.read_bytes()
            with patch.object(self.support.json, "dump", side_effect=OSError("simulated disk failure")) as dump:
                self.assertFalse(self.support.persist_reverse_outline_result({"status": "ok", "tree": tree}, data_dir))
                dump.assert_called_once()
            self.assertEqual(target.read_bytes(), old_bytes)
            self.assertEqual({path.name for path in Path(data_dir).iterdir()}, {"logic_tree.json"})

    def test_both_app_reverse_tabs_have_no_fixed_truncation_and_save_only_success(self):
        support = self.support

        class UI:
            def __init__(self, document_text="complete source"):
                self.messages = []
                self.document_text = document_text
            def __getattr__(self, name):
                if name == "file_uploader":
                    return lambda *args, **kwargs: types.SimpleNamespace(read=lambda: b"docx")
                if name == "button":
                    return lambda *args, **kwargs: True
                if name == "spinner":
                    return lambda message: nullcontext()
                if name in ("subheader", "info", "success", "warning", "rerun"):
                    return lambda *args, **kwargs: self.messages.append((name, args))
                if name == "error":
                    return lambda *args, **kwargs: self.messages.append((name, args))
                raise AssertionError(f"Unexpected Streamlit UI call: {name}")

        for app_name, locale in (("app.py", "en"), ("app_zh.py", "zh")):
            with self.subTest(app=app_name):
                source = (ROOT / app_name).read_text(encoding="utf-8")
                app_module = ast.parse(source)
                reverse_tab = next(node for node in app_module.body if isinstance(node, ast.FunctionDef) and node.name == "ast_reverse_tab")
                reverse_tab_source = ast.get_source_segment(source, reverse_tab)
                self.assertNotIn("extracted_text[:4000]", reverse_tab_source)
                ui = UI()
                with tempfile.TemporaryDirectory() as data_dir:
                    namespace = {
                        "st": ui,
                        "io": io,
                        "docx": types.SimpleNamespace(Document=lambda _file, text=ui.document_text: types.SimpleNamespace(paragraphs=[types.SimpleNamespace(text=text)] if text else [])),
                        "extract_docx_manuscript_text": lambda doc: "\n".join(p.text for p in doc.paragraphs),
                        "reverse_engineer_manuscript": lambda text, llm_call, locale: {"status": "error", "error_type": "parse"},
                        "persist_reverse_outline_result": support.persist_reverse_outline_result,
                        "DATA_DIR": data_dir,
                        "dispatch_llm_call": lambda *args, **kwargs: self.fail("failure result should not call LLM"),
                        "DEFAULT_CHUNK_CHARS": self.documents.DEFAULT_CHUNK_CHARS,
                        "llm_error_message": lambda _locale: "model configuration message",
                    }
                    function = load_app_function(self, app_name, "ast_reverse_tab", namespace)
                    function()
                    self.assertFalse((Path(data_dir) / "logic_tree.json").exists())
                    self.assertTrue(any(name == "error" for name, _ in ui.messages))

                    empty_calls = []
                    empty_ui = UI("")
                    empty_namespace = dict(namespace)
                    empty_namespace.update({
                        "st": empty_ui,
                        "docx": types.SimpleNamespace(Document=lambda _file: types.SimpleNamespace(paragraphs=[])),
                        "extract_docx_manuscript_text": lambda _doc: "",
                        "reverse_engineer_manuscript": lambda *args, **kwargs: empty_calls.append(args),
                    })
                    load_app_function(self, app_name, "ast_reverse_tab", empty_namespace)()
                    self.assertEqual(empty_calls, [])
                    self.assertFalse((Path(data_dir) / "logic_tree.json").exists())


class ChartSupportTests(unittest.TestCase):
    def setUp(self):
        self.chart = load_project_module(self, "chart_support")

    def load_chart_support(self):
        return self.chart

    def parse(self, spec, instruction="A=10, B=20", **kwargs):
        chart = self.load_chart_support()
        return chart.parse_chart_spec(json.dumps(spec), instruction=instruction, **kwargs)

    def test_valid_grouped_bar_chart_spec_is_accepted(self):
        spec = self.parse({
            "status": "ok",
            "chart_type": "bar",
            "title": "Comparison",
            "x_label": "Category",
            "y_label": "Value",
            "categories": ["A", "B"],
            "series": [
                {"name": "Group 1", "values": [10, 20]},
                {"name": "Group 2", "values": [5, 8]},
            ],
        }, instruction="A=10, B=20; Group 2 has A=5 and B=8")
        self.assertEqual(spec["status"], "ok")
        self.assertEqual(spec["chart_type"], "bar")
        self.assertEqual(len(spec["series"]), 2)

    def test_bar_series_length_must_match_categories(self):
        chart = self.load_chart_support()
        with self.assertRaises(chart.ChartSpecError):
            self.parse({
                "status": "ok", "chart_type": "bar", "title": "T",
                "x_label": "X", "y_label": "Y", "categories": ["A", "B"],
                "series": [{"name": "S", "values": [10]}],
            })

    def test_valid_line_chart_requires_matching_x_and_series_values(self):
        spec = self.parse({
            "status": "ok", "chart_type": "line", "title": "Trend",
            "x_label": "Year", "y_label": "Value", "x": ["2020", "2021"],
            "series": [{"name": "A", "values": [10, 20]}],
        }, instruction="In 2020 value was 10; in 2021 value was 20")
        self.assertEqual(spec["chart_type"], "line")

    def test_line_series_length_mismatch_is_rejected(self):
        chart = self.load_chart_support()
        with self.assertRaises(chart.ChartSpecError):
            self.parse({
                "status": "ok", "chart_type": "line", "title": "T",
                "x_label": "X", "y_label": "Y", "x": ["2020", "2021"],
                "series": [{"name": "S", "values": [10]}],
            }, instruction="2020=10, 2021=20")

    def test_scatter_x_and_y_length_mismatch_is_rejected(self):
        chart = self.load_chart_support()
        with self.assertRaises(chart.ChartSpecError):
            self.parse({
                "status": "ok", "chart_type": "scatter", "title": "T",
                "x_label": "X", "y_label": "Y",
                "series": [{"name": "S", "x": [1, 2], "y": [3]}],
            }, instruction="(1,3), (2,4)")

    def test_pie_rejects_negative_values_and_zero_total(self):
        chart = self.load_chart_support()
        base = {"status": "ok", "chart_type": "pie", "title": "T", "labels": ["A", "B"]}
        for values in ([-1, 2], [0, 0]):
            with self.subTest(values=values), self.assertRaises(chart.ChartSpecError):
                self.parse({**base, "values": values}, instruction="A=-1, B=2")

    def test_non_finite_json_numbers_are_rejected(self):
        chart = self.load_chart_support()
        for token in ("NaN", "Infinity", "-Infinity"):
            raw = (
                '{"status":"ok","chart_type":"bar","title":"T",'
                '"x_label":"X","y_label":"Y","categories":["A"],'
                f'"series":[{{"name":"S","values":[{token}]}}]}}'
            )
            with self.subTest(token=token), self.assertRaises(chart.ChartSpecError):
                chart.parse_chart_spec(raw, instruction="A=10")

    def test_boolean_is_not_accepted_as_a_number(self):
        chart = self.load_chart_support()
        with self.assertRaises(chart.ChartSpecError):
            self.parse({
                "status": "ok", "chart_type": "bar", "title": "T",
                "x_label": "X", "y_label": "Y", "categories": ["A"],
                "series": [{"name": "S", "values": [True]}],
            }, instruction="A=1")

    def test_unsupported_chart_types_and_llm_code_fields_are_rejected(self):
        chart = self.load_chart_support()
        valid = {
            "status": "ok", "chart_type": "python", "title": "T",
            "x_label": "X", "y_label": "Y", "categories": ["A"],
            "series": [{"name": "S", "values": [10]}],
        }
        for unsupported in ("python", [], {}):
            with self.subTest(chart_type=unsupported), self.assertRaises(chart.ChartSpecError):
                self.parse({**valid, "chart_type": unsupported})
        valid["chart_type"] = "bar"
        valid["code"] = "import os; os.remove('data')"
        with self.assertRaises(chart.ChartSpecError):
            self.parse(valid)

    def test_needs_data_is_returned_as_a_non_renderable_result(self):
        chart = self.load_chart_support()
        result = chart.parse_chart_spec(
            '{"status":"needs_data","message":"Please provide values."}',
            instruction="Compare the three groups",
        )
        self.assertEqual(result["status"], "needs_data")

    def test_model_needs_data_result_is_translated_without_rendering(self):
        chart = self.load_chart_support()
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "chart.png"
            result = chart.generate_chart_image(
                "Compare A=10 and B=20",
                str(output),
                lambda *args, **kwargs: '{"status":"needs_data","message":"Need another value."}',
            )
            self.assertEqual(result["status"], "needs_data")
            self.assertIn("More data is needed", result["message"])
            self.assertFalse(output.exists())

    def test_invalid_spec_retry_uses_json_mode_and_validation_feedback(self):
        chart = self.load_chart_support()
        responses = iter((
            '{"status":"ok","chart_type":"unsupported"}',
            '{"status":"needs_data","message":"Please provide values."}',
        ))
        calls = []

        def llm_call(prompt, **kwargs):
            calls.append((prompt, kwargs))
            return next(responses)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "chart.png"
            result = chart.generate_chart_image("Compare A=10 and B=20", str(output), llm_call)
            self.assertEqual(result["status"], "needs_data")
            self.assertEqual(len(calls), 2)
            self.assertTrue(all(call[1]["json_mode"] for call in calls))
            self.assertIn("previous JSON was invalid", calls[1][0])
            self.assertFalse(output.exists())

    def test_factual_values_must_come_from_the_user_instruction(self):
        chart = self.load_chart_support()
        spec = {
            "status": "ok", "chart_type": "bar", "title": "T",
            "x_label": "X", "y_label": "Y", "categories": ["A", "B"],
            "series": [{"name": "S", "values": [10, 99]}],
        }
        with self.assertRaises(chart.ChartSpecError):
            chart.parse_chart_spec(json.dumps(spec), instruction="A=10, B=20")

    def test_line_numeric_axis_labels_must_come_from_the_user_instruction(self):
        chart = self.load_chart_support()
        spec = {
            "status": "ok", "chart_type": "line", "title": "Trend",
            "x_label": "Year", "y_label": "Value", "x": ["1990", "1991"],
            "series": [{"name": "S", "values": [10, 20]}],
        }
        with self.assertRaises(chart.ChartSpecError):
            chart.parse_chart_spec(json.dumps(spec), instruction="2020=10, 2021=20")

    def test_illustrative_request_is_allowed_and_marked(self):
        chart = self.load_chart_support()
        spec = {
            "status": "ok", "chart_type": "bar", "title": "Example",
            "x_label": "X", "y_label": "Y", "categories": ["A", "B"],
            "series": [{"name": "S", "values": [10, 20]}],
        }
        result = chart.parse_chart_spec(
            json.dumps(spec), instruction="Use hypothetical illustrative data for A and B"
        )
        self.assertTrue(result["data_note"])

    def test_chinese_illustrative_permission_requires_explicit_data_authorization(self):
        chart = self.load_chart_support()
        permitted = (
            "请使用示例数据绘图",
            "可以生成一组假设数据",
            "允许使用示意数值",
            "没有真实数据，可以自行生成演示数据",
            "示例数据即可",
            "假设数据即可",
        )
        for instruction in permitted:
            with self.subTest(instruction=instruction):
                self.assertTrue(chart.is_illustrative_request(instruction))

    def test_negative_illustrative_wording_does_not_allow_fabricated_values(self):
        chart = self.load_chart_support()
        for instruction in (
            "不要使用示例数据",
            "不使用示例数据",
            "不采用假设数据",
            "别用示意数据",
            "请勿使用示例数据",
            "禁止使用演示数据",
            "不能生成假设数字",
            "不得采用示例数值",
            "避免使用假设数据",
            "不可生成假设数字",
            "Compare values, but do not use illustrative data",
            "No hypothetical numbers; use only my research data",
            "不要使用示例数据，请使用我提供的数值",
        ):
            with self.subTest(instruction=instruction):
                self.assertFalse(chart.is_illustrative_request(instruction))

    def test_example_and_hypothetical_mentions_without_data_permission_are_not_allowed(self):
        chart = self.load_chart_support()
        mentions = (
            "参考示例图的样式绘制",
            "按照示例格式绘制",
            "这是一个假设场景，但请使用我的真实数据",
            "请仿照示例图布局",
            "示例格式如下",
            "讨论一个假设案例",
            "Follow the example chart's style.",
            "Use the sample figure layout.",
            "This is a hypothetical scenario, but use my real data.",
            "Use the example formatting.",
        )
        for instruction in mentions:
            with self.subTest(instruction=instruction):
                self.assertFalse(chart.is_illustrative_request(instruction))

    def test_english_illustrative_permission_and_negation_semantics(self):
        chart = self.load_chart_support()
        permitted = (
            "Use illustrative data.",
            "Generate hypothetical values for this example.",
            "Sample data is fine.",
            "You may make up example values for demonstration.",
            "Use hypothetical numbers.",
        )
        denied = (
            "Do not use illustrative data.",
            "Don't generate hypothetical values.",
            "No sample data.",
            "Without example values.",
            "Never make up illustrative numbers.",
        )
        for instruction in permitted:
            with self.subTest(permitted=instruction):
                self.assertTrue(chart.is_illustrative_request(instruction))
        for instruction in denied:
            with self.subTest(denied=instruction):
                self.assertFalse(chart.is_illustrative_request(instruction))

    def test_separate_explicit_permission_clause_is_still_allowed(self):
        chart = self.load_chart_support()
        self.assertTrue(chart.is_illustrative_request(
            "不要用第一组示例数据；第二张图可以使用假设数据。"
        ))

    def test_negated_illustrative_request_needs_data_without_calling_llm_or_creating_png(self):
        chart = self.load_chart_support()
        calls = []

        def unexpected_llm_call(*args, **kwargs):
            calls.append((args, kwargs))
            return "must not be called"

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "chart.png"
            result = chart.generate_chart_image(
                "不使用示例数据，请画三个组的比较图", str(output), unexpected_llm_call, locale="zh"
            )
            self.assertEqual(result["status"], "needs_data")
            self.assertFalse(calls)
            self.assertFalse(output.exists())

    def test_explicit_illustrative_permission_enters_chart_spec_flow(self):
        chart = self.load_chart_support()
        calls = []
        response = json.dumps({
            "status": "ok", "chart_type": "bar", "title": "演示比较",
            "x_label": "组别", "y_label": "数值", "categories": ["A", "B", "C"],
            "series": [{"name": "示意组", "values": [10, 20, 30]}],
        })

        def llm_call(prompt, **kwargs):
            calls.append((prompt, kwargs))
            return response

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "chart.png"
            result = chart.generate_chart_image(
                "请使用假设数据画三个组的演示柱状图", str(output), llm_call, locale="zh"
            )
            self.assertEqual(result["status"], "ok")
            self.assertEqual(len(calls), 1)
            self.assertTrue(calls[0][1]["json_mode"])
            self.assertEqual(result["spec"]["data_note"], "示例数据（假设）")
            self.assertTrue(chart.chart_png_is_valid(str(output)))

    def test_year_only_request_returns_needs_data_without_calling_llm(self):
        chart = self.load_chart_support()
        no_data_requests = (
            "Plot the trend from 2020 to 2024",
            "Use data from 2020 to 2024",
            "Draw a 3D chart with 2 series",
        )
        self.assertTrue(chart.has_usable_numeric_data("The average score was 3.14"))
        self.assertTrue(chart.has_usable_numeric_data("A=2024"))
        for index, instruction in enumerate(no_data_requests):
            with self.subTest(instruction=instruction), tempfile.TemporaryDirectory() as tmp:
                output = Path(tmp) / f"chart-{index}.png"
                calls = []

                def unexpected_llm_call(*args, **kwargs):
                    calls.append((args, kwargs))
                    return "should not be called"

                result = chart.generate_chart_image(instruction, str(output), unexpected_llm_call)
                self.assertEqual(result["status"], "needs_data")
                self.assertFalse(calls)
                self.assertFalse(output.exists())

    def test_prompt_requests_json_and_forbids_invented_factual_values(self):
        chart = self.load_chart_support()
        prompt = chart.build_chart_spec_prompt("Compare A=10 and B=20", locale="en")
        self.assertIn("bar", prompt)
        self.assertIn("scatter", prompt)
        self.assertIn("pie", prompt)
        self.assertIn("needs_data", prompt)
        self.assertIn("A=10 and B=20", prompt)
        self.assertRegex(prompt.lower(), r"never (?:invent|infer)")
        self.assertNotIn("python script", prompt.lower())
        self.assertNotIn("python code", prompt.lower())
        self.assertNotIn("Output only pure Python code", prompt)

    def test_renderer_writes_valid_bar_and_line_pngs(self):
        chart = self.load_chart_support()
        fixtures = [
            (
                "bar", "A=10 and B=20",
                {"status": "ok", "chart_type": "bar", "title": "Bars",
                 "x_label": "Category", "y_label": "Value", "categories": ["A", "B"],
                 "series": [{"name": "S", "values": [10, 20]}]},
            ),
            (
                "line", "2020=10 and 2021=20",
                {"status": "ok", "chart_type": "line", "title": "Trend",
                 "x_label": "Year", "y_label": "Value", "x": ["2020", "2021"],
                 "series": [{"name": "S", "values": [10, 20]}]},
            ),
            (
                "scatter", "x=1 y=4, x=2 y=6",
                {"status": "ok", "chart_type": "scatter", "title": "Relationship",
                 "x_label": "X", "y_label": "Y",
                 "series": [{"name": "S", "x": [1, 2], "y": [4, 6]}]},
            ),
            (
                "pie", "A=30, B=45, C=25",
                {"status": "ok", "chart_type": "pie", "title": "Composition",
                 "labels": ["A", "B", "C"], "values": [30, 45, 25]},
            ),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            for name, instruction, raw_spec in fixtures:
                with self.subTest(chart=name):
                    spec = chart.parse_chart_spec(json.dumps(raw_spec), instruction=instruction)
                    path = Path(tmp) / f"{name}.png"
                    chart.render_chart_spec(spec, str(path))
                    payload = path.read_bytes()
                    self.assertTrue(payload.startswith(b"\x89PNG\r\n\x1a\n"))
                    self.assertGreater(len(payload), 32)
                    width, height = struct.unpack(">II", payload[16:24])
                    self.assertGreater(width, 0)
                    self.assertGreater(height, 0)
                    self.assertLessEqual(width * height, chart.MAX_CHART_PIXELS)

    def test_illustrative_chart_is_rendered_with_a_visible_data_note(self):
        chart = self.load_chart_support()
        raw_spec = {
            "status": "ok", "chart_type": "bar", "title": "Illustration",
            "x_label": "Category", "y_label": "Value", "categories": ["A", "B"],
            "series": [{"name": "S", "values": [10, 20]}],
        }
        spec = chart.parse_chart_spec(
            json.dumps(raw_spec), instruction="Use hypothetical illustrative data for A and B"
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "illustrative.png"
            chart.render_chart_spec(spec, str(path))
            self.assertTrue(chart.chart_png_is_valid(str(path)))

    def test_llm_failure_and_invalid_specs_do_not_create_chart_files(self):
        chart = self.load_chart_support()
        with tempfile.TemporaryDirectory() as tmp:
            for response, expected_status in (
                ("API call failed repeatedly: timeout", "error"),
                ('{"status":"ok","chart_type":"python"}', "error"),
            ):
                with self.subTest(expected_status=expected_status), tempfile.TemporaryDirectory(dir=tmp) as case_dir:
                    output = Path(case_dir) / "chart.png"
                    result = chart.generate_chart_image(
                        "Compare A=10 and B=20", str(output), lambda *args, **kwargs: response
                    )
                    self.assertEqual(result["status"], expected_status)
                    self.assertFalse(output.exists())

    def test_both_apps_have_no_llm_python_execution_path(self):
        for name in ("app.py", "app_zh.py"):
            with self.subTest(app=name):
                source = (ROOT / name).read_text(encoding="utf-8")
                self.assertFalse("_clean_python_code" in source, "legacy executable chart cleaner remains")
                self.assertFalse(re.search(r"\bcompile\s*\(\s*cleaned_code", source), "chart code is still compiled")
                self.assertFalse(re.search(r"\bexec\s*\(\s*cleaned_code", source), "chart code is still executed")
                self.assertTrue("generate_chart_image" in source, "shared structured chart flow is not used")
                self.assertFalse(re.search(r"(?:Output only|只输出).{0,30}Python code", source), "chart prompt still requests Python")
                chart_flow = source.split("def generate_chart_for_node", 1)[1].split("def build_chapter_prompt", 1)[0]
                self.assertNotRegex(chart_flow.lower(), r"\bpython\b", "chart generation still contains a Python-code instruction")
                self.assertLess(
                    chart_flow.index('if result["status"] != "ok":'),
                    chart_flow.index('load_json_file("drafts_charts.json", {})'),
                    "chart persistence must happen only after successful validation and rendering",
                )
                self.assertIn("generate_chart_for_node(nid, instruction)", source, "batch chart generation bypasses the shared safe flow")
        chart_source = (ROOT / "chart_support.py").read_text(encoding="utf-8")
        self.assertTrue("json_mode=True" in chart_source, "chart LLM call does not request JSON mode")
        self.assertFalse(re.search(r"(?<![.\w])(?:exec|eval|compile)\s*\(", chart_source), "chart helper executes code")


if __name__ == "__main__":
    unittest.main()
